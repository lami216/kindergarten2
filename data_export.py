"""Verified, lossless data export. A data archive, not a database restore image."""
import base64
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
import io
import json
import zipfile

from openpyxl import Workbook
from psycopg2 import sql
import db


def _json(value):
    if isinstance(value, (date, datetime, Decimal)):
        return str(value)
    if isinstance(value, (bytes, memoryview)):
        return {"base64": base64.b64encode(bytes(value)).decode("ascii")}
    raise TypeError(f"Unsupported export type: {type(value).__name__}")


def _bytes(value):
    return json.dumps(value, default=_json, ensure_ascii=False, sort_keys=True).encode("utf-8")


def _cell(value):
    if value is None or isinstance(value, (bool, int, float)):
        return value
    text = json.dumps(value, default=_json, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)
    # Excel's cell limit must never truncate the authoritative JSON archive.
    if len(text) > 32000 or any(ord(c) < 32 and c not in '\t\n\r' for c in text):
        return "القيمة الكاملة محفوظة في ملف JSON لهذا الجدول"
    # Treat user text as text, never an executable spreadsheet formula.
    return "'" + text if text.startswith(('=', '+', '-', '@')) else text


def build_export():
    connection = db.new_transaction_connection()
    try:
        # The factory sets audit context; this read-only snapshot starts afresh.
        connection.rollback()
        connection.set_session(isolation_level="REPEATABLE READ", readonly=True)
        workbook = Workbook()
        index = workbook.active
        index.title = 'فهرس التصدير'
        index.append(['الجدول', 'التبويب', 'عدد السجلات', 'ملف البيانات الكامل', 'SHA256'])
        files = {}
        manifest = {'format': 1, 'created_at': datetime.now(timezone.utc).isoformat(),
                    'purpose': 'Complete table data export; not a database restore image', 'tables': []}
        with connection.cursor() as cursor:
            import access_control as A
            A.require(cursor,'admin.export')
            cursor.execute("SELECT table_name FROM information_schema.tables WHERE table_schema=current_schema() AND table_type='BASE TABLE' ORDER BY table_name")
            tables = [r['table_name'] for r in cursor.fetchall()]
            if not tables:
                raise ValueError('لا توجد جداول للتصدير.')
            for number, table in enumerate(tables, 1):
                cursor.execute("SELECT column_name,data_type,udt_name,is_nullable,column_default FROM information_schema.columns WHERE table_schema=current_schema() AND table_name=%s ORDER BY ordinal_position", (table,))
                columns = [dict(r) for r in cursor.fetchall()]
                cursor.execute(sql.SQL('SELECT * FROM {}').format(sql.Identifier(table)))
                rows = [dict(r) for r in cursor.fetchall()]
                name = 'data/' + table + '.json'
                payload = _bytes(rows)
                files[name] = payload
                sheet_name = {'monthly_receipts': 'الدفعات الشهرية', 'payments': 'دفعات سابقة',
                              'record_changes': 'سجل التغييرات'}.get(table, f'{number:02d}_{table}'[:31])
                sheet = workbook.create_sheet(sheet_name)
                keys = [r['column_name'] for r in columns]
                sheet.append(keys)
                for row in rows:
                    sheet.append([_cell(row[key]) for key in keys])
                digest = hashlib.sha256(payload).hexdigest()
                manifest['tables'].append({'name': table, 'rows': len(rows), 'columns': columns,
                                           'file': name, 'sha256': digest, 'sheet': sheet_name})
                index.append([table, sheet_name, len(rows), name, digest])
        excel = io.BytesIO()
        workbook.save(excel)
        files['records.xlsx'] = excel.getvalue()
        files['manifest.json'] = _bytes(manifest)
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, payload in files.items():
                archive.writestr(name, payload)
        result = output.getvalue()
        verify_export(result)
        connection.commit()
        return result, manifest
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def verify_export(payload):
    """Refuse a partial archive, a mismatched row count or changed JSON content."""
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        if archive.testzip() is not None:
            raise ValueError('ملف التصدير غير مكتمل.')
        manifest = json.loads(archive.read('manifest.json'))
        for table in manifest['tables']:
            content = archive.read(table['file'])
            if hashlib.sha256(content).hexdigest() != table['sha256'] or len(json.loads(content)) != table['rows']:
                raise ValueError('تعذّر التحقق من اكتمال بيانات التصدير.')
        return manifest
