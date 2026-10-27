"""Period student totals, permission boundaries and shared-family corrections."""
from datetime import timedelta
from io import BytesIO
import uuid

from openpyxl import load_workbook
import access_control as A
import monthly_finance as F
from tests.test_access_control import access, PASSWORD
from tests.test_live_pages_http import browser, intake, login, post


def period(client, session, start, end):
    response = post(client, session, 'reports.finance', dict(start=str(start), end=str(end)))
    assert response.status_code == 200, response.text
    return response.json()


def child(client, session, **changes):
    response = post(client, session, 'students.create', intake(**changes))
    assert response.status_code == 200, response.text
    return response.json()['id']


def get_student(client, ident):
    return next(s for s in client.get('/api/snapshot').json()['students'] if s['id'] == ident)


def test_period_enrollment_dates_archives_and_weighted_attendance(browser):
    client, conn = browser; session = login(client)
    today = F.today(); start = today - timedelta(days=2)
    first = child(client, session, name='Same name')
    second = child(client, session, name='Same name')
    with conn.cursor() as c:
        c.execute('UPDATE students SET created_at=%s WHERE student_id=%s', (start, first))
        c.execute('UPDATE students SET created_at=%s WHERE student_id=%s', (today, second))
        # An older real registration takes precedence over a later file creation.
        c.execute('INSERT INTO registrations(student_id,status,registration_date) VALUES(%s,%s,%s)',
                  (second, 'نشط', start-timedelta(days=1)))
    entries = [(first, start, 'present', ''), (first, start+timedelta(days=1), 'absent', ''),
               (first, today, 'excused', 'Medical visit'), (second, today, 'present', '')]
    for ident, day, status, excuse in entries:
        r = post(client, session, 'attendance.save', dict(kind='students', date=str(day),
            entries=[dict(person=ident, status=status, excuse=excuse, revision=0)]))
        assert r.status_code == 200, r.text
    # Cancelled records do not turn into an absence; historic records survive archive.
    with conn.cursor() as c:
        c.execute('UPDATE students SET archived_at=now() WHERE student_id=%s', (first,))
    report = period(client, session, start, today)
    school = report['school']; stats = school['attendance']
    assert school['new_count'] == 1 and school['new_students'][0]['id'] == first
    assert school['new_students'][0]['archived'] is True
    assert stats['summary'] == dict(present=2, absent=1, excused=1, total=4, attendance_rate=50.0, absence_rate=50.0)
    assert stats['recorded_days'] == 3 and stats['students_count'] == 2
    assert len({row['id'] for row in stats['items']}) == 2
    exported = post(client, session, 'reports.finance.export', dict(start=str(start), end=str(today), revision=report['revision']))
    assert exported.status_code == 200, exported.text
    book = load_workbook(BytesIO(exported.content))
    assert book['حضور الطلاب'].max_row == 3 and book['الطلاب الجدد'].max_row == 2
    empty = period(client, session, start-timedelta(days=3), start-timedelta(days=2))['school']['attendance']
    assert empty['summary']['total'] == 0 and empty['summary']['attendance_rate'] is None


def test_report_redaction_and_attendance_correction_invalidate_export(browser):
    client, conn = browser; session = login(client); today = F.today()
    sid = child(client, session, name='Private school child')
    r = post(client, session, 'attendance.save', dict(kind='students', date=str(today), entries=[dict(person=sid,status='absent',excuse='',revision=0)]))
    assert r.status_code == 200, r.text
    original = period(client, session, today, today)
    day = post(client, session, 'attendance.read', dict(kind='students', date=str(today))).json()
    entry = next(row['entry'] for row in day['items'] if row['id'] == sid)
    r = post(client, session, 'attendance.save', dict(kind='students', date=str(today), entries=[dict(person=sid,status='',excuse='',revision=entry['revision'])]))
    assert r.status_code == 200, r.text
    changed = period(client, session, today, today)
    assert changed['school']['attendance']['summary']['total'] == 0
    assert changed['revision'] != original['revision']
    assert post(client,session,'reports.finance.export',dict(start=str(today),end=str(today),revision=original['revision'])).status_code == 422
    A.save_user(employee_name='Reports only',name='school_report_only',password=PASSWORD,permissions=['reports.view'])
    scoped = login(client, 'school_report_only')
    hidden = period(client, scoped, today, today)
    assert hidden['school']['new_count'] == 1
    assert hidden['school']['new_students'] == [] and hidden['school']['attendance'] is None
    assert 'Private school child' not in str(hidden)
    export = post(client,scoped,'reports.finance.export',dict(start=str(today),end=str(today),revision=hidden['revision']))
    book=load_workbook(BytesIO(export.content))
    assert 'الطلاب الجدد' not in book.sheetnames and 'حضور الطلاب' not in book.sheetnames


def test_shared_parent_correction_is_atomic_audited_and_conflict_protected(browser):
    client, conn = browser; session = login(client)
    one = child(client, session)
    first = get_student(client, one)
    two = child(client, session, name='Sibling', existing_parent_id=first['parent_id'])
    stale = get_student(client, two)
    r = post(client, session, 'students.edit', dict(id=one,revision=first['revision'],parent_revision=first['parent_revision'],
        reason='Correct parent phone',values=dict(motherPhone='99887766',motherName='Corrected mother',address='New address')))
    assert r.status_code == 200, r.text
    assert get_student(client, two)['motherPhone'] == '99887766'
    assert get_student(client, two)['parent_revision'] != stale['parent_revision']
    rejected = post(client, session, 'students.edit', dict(id=two,revision=stale['revision'],parent_revision=stale['parent_revision'],
        reason='Stale sibling form',values=dict(motherPhone='11111111',address='Must roll back')))
    assert rejected.status_code == 422
    current = get_student(client, two)
    assert current['address'] != 'Must roll back' and current['motherPhone'] == '99887766'
    invalid = post(client, session, 'students.edit',dict(id=two,revision=current['revision'],parent_revision=current['parent_revision'],
        reason='Invalid pair',values=dict(motherPhone='',address='Invalid rollback')))
    assert invalid.status_code == 422 and get_student(client, two)['address'] != 'Invalid rollback'
    with conn.cursor() as c:
        c.execute("SELECT actor FROM record_changes WHERE table_name='parents' AND record_key=%s AND action='updated'", (first['parent_id'],))
        assert any(row['actor']=='owner' for row in c.fetchall())


def test_student_editor_cannot_change_parents_without_parent_permission(browser):
    client, conn = browser; owner = login(client); sid = child(client, owner)
    A.save_user(employee_name='Student editor',name='child_editor',password=PASSWORD,permissions=['students.view','students.edit'])
    session = login(client, 'child_editor'); student = get_student(client, sid)
    r = post(client,session,'students.edit',dict(id=sid,revision=student['revision'],parent_revision=student['parent_revision'],
        reason='Forbidden family edit',values=dict(address='Not saved',motherPhone='99887766')))
    assert r.status_code == 403
    current=get_student(client,sid)
    assert current['address'] != 'Not saved'
    r=post(client,session,'students.edit',dict(id=sid,revision=current['revision'],reason='Student only',values=dict(address='Allowed address')))
    assert r.status_code == 200, r.text
