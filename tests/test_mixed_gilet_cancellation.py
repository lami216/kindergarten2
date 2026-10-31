"""Mixed-receipt corrections, exact retained allocations and real HTTP guards."""
from datetime import timedelta
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
import uuid
import pytest
from openpyxl import load_workbook
import monthly_finance as F
import access_control as A
from tests.test_access_control import access,PASSWORD
from tests.test_live_pages_http import browser,login,post,intake,ORIGIN


def create(client,session,**values):
    r=post(client,session,'students.create',intake(mode='pay',registration=True,gilet=True,amount='6000',**values))
    assert r.status_code==200,r.text
    state=client.get('/api/snapshot').json()
    sid=r.json()['id'];return sid,next(r['id'] for r in state['receipts'] if r['student']==sid)


def plan(client,session,rid):
    r=post(client,session,'receipts.gilet.preview',dict(id=rid))
    assert r.status_code==200,r.text
    return r.json()


def execute(client,session,rid,**overrides):
    p=plan(client,session,rid)
    data=dict(id=rid,revision=p['revision'],reason='Correct garment item only',confirmed=True,request_token=str(uuid.uuid4()))
    data.update(overrides)
    return post(client,session,'receipts.gilet.cancel',data),data


def allocations(conn,rid):
    with conn.cursor() as c:
        c.execute('''SELECT a.*,s.fee_kind,s.period_no,s.starts,s.ends,s.stored_fee_price,s.source
            FROM monthly_allocations a JOIN monthly_allocation_snapshots s USING(allocation_id)
            WHERE receipt_id=%s ORDER BY allocation_id''',(rid,))
        return [dict(r) for r in c.fetchall()]


@pytest.mark.parametrize('monthly',[False,True])
def test_cancel_only_garment_keeps_registration_and_months_and_reports(browser,monthly):
    client,conn=browser;session=login(client)
    # The tariff is a test fixture, not hard-coded application pricing.
    data=intake(mode='pay',registration=True,gilet=True,amount='6000' if monthly else '3000')
    response=post(client,session,'students.create',data);assert response.status_code==200,response.text
    sid=response.json()['id'];before=F.projection(conn,sid)
    rid=before['receipts'][0]['receipt_id'];original=allocations(conn,rid)
    p=plan(client,session,rid)
    assert p['eligible'] and p['gilet']==2000 and p['registration']==1000 and p['monthly']==(3000 if monthly else 0)
    r,payload=execute(client,session,rid);assert r.status_code==200,r.text
    new=r.json()['id'];assert new!=rid
    after=F.projection(conn,sid);kept=allocations(conn,new)
    assert after['cash_paid']==(4000 if monthly else 1000) and after['due_now']==before['due_now']
    assert after['fees']['gilet']['cancelled']==2000
    assert after['periods']==before['periods'] and after['account']==before['account']
    correct=[x for x in original if x['fee_kind']!='gilet']
    keys=('fee_id','period_id','gilet_sale_id','amount','fee_kind','period_no','starts','ends','stored_fee_price','source')
    assert [{k:x[k] for k in keys} for x in kept]==[{k:x[k] for k in keys} for x in correct]
    assert all(x['copied_from_id']==source['allocation_id'] for x,source in zip(kept,correct))
    old_after=allocations(conn,rid)
    assert [{k:v for k,v in x.items() if k!='active'} for x in old_after]==[{k:v for k,v in x.items() if k!='active'} for x in original]
    assert not any(x['active'] for x in old_after)
    replay=post(client,session,'receipts.gilet.cancel',payload)
    assert replay.status_code==200 and replay.json()['id']==new and replay.json()['created'] is False
    assert post(client,session,'receipts.gilet.cancel',{**payload,'reason':'Different request'}).status_code==422
    report=post(client,session,'reports.finance',dict(start=str(F.today()),end=str(F.today()))).json()
    assert report['summary']['gilet']==0 and report['summary']['registration']==1000
    assert report['summary']['monthly']==(3000 if monthly else 0)
    assert report['summary']['income']==(4000 if monthly else 1000)
    summary=F.financial_report(conn,F.today(),F.today());assert summary['gilet_cash']==0 and summary['total_cash']==report['summary']['income']
    exported=post(client,session,'reports.finance.export',dict(start=str(F.today()),end=str(F.today()),revision=report['revision']))
    assert exported.status_code==200 and len(load_workbook(BytesIO(exported.content)).sheetnames)==6
    for ident in (rid,new):
        pdf=client.get('/api/receipt/'+str(ident));assert pdf.status_code==200 and pdf.content.startswith(b'%PDF')
    with conn.cursor() as c:
        c.execute('SELECT actor,reason,amount FROM monthly_gilet_cancellations WHERE receipt_id=%s',(rid,))
        audit=c.fetchone();assert audit['actor']=='owner' and audit['reason']==payload['reason'] and audit['amount']==2000
    assert not plan(client,session,new)['eligible'] and not plan(client,session,rid)['eligible']


def test_tariff_anchor_and_receipt_date_changes_do_not_move_retained_history(browser):
    from account_records import change_account
    from receipt_history import receipt_detail
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    original=allocations(conn,rid);anchor=F.today()-timedelta(days=40)
    F.update_tariffs(1500,2800,4500,'owner')
    change_account(sid,actor='owner',reason='Correct subscription date',start_date=anchor,status='regular')
    prior=F.projection(conn,sid)
    result,_=execute(client,session,rid);assert result.status_code==200,result.text
    current=result.json()['id'];after=F.projection(conn,sid)
    assert after['periods']==prior['periods'] and after['due_now']==prior['due_now']
    assert after['account']['start_date']==anchor
    expected=next(x for x in original if x['period_id'])
    assert next(x for x in allocations(conn,current) if x['period_id'])['starts']==expected['starts']
    corrected,_=F.correct_receipt(current,4000,'Payment date only',received_on=F.today()-timedelta(days=2),actor='owner')
    actual=next(x for x in allocations(conn,corrected) if x['period_id'])
    assert actual['starts']==expected['starts'] and actual['period_id']==expected['period_id'] and actual['stored_fee_price']==3000
    assert receipt_detail(conn,corrected)['receipt']['amount']==4000
    assert F.projection(conn,sid)['due_now']==prior['due_now']


def test_other_garments_delivery_and_new_purchase_remain_independent(browser):
    from tests.test_gilet_exemption import sale
    import gilet_sales as G
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    p=G.preview(conn,sid);old=p['purchases'][0]
    G.deliver(sid,old['id'],actor='owner',reason='Confirmed actual delivery',confirmed=True,token=str(uuid.uuid4()),expected_revision=p['ledger_revision'])
    independent,_=sale(client,session,sid);assert independent.status_code==200
    result,_=execute(client,session,rid);assert result.status_code==200,result.text
    items=G.preview(conn,sid)['purchases'];legacy=next(x for x in items if x['kind']=='legacy')
    assert legacy['delivered'] and legacy['cancelled'] and legacy['due']==0
    assert next(x for x in items if x['kind']=='sale')['paid']==2000
    again,_=sale(client,session,sid);assert again.status_code==200
    assert F.projection(conn,sid)['cash_paid']==8000


def test_permission_origin_csrf_and_midtransaction_failure_leave_no_partial_effect(browser,monkeypatch):
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    p=plan(client,session,rid);payload=dict(id=rid,revision=p['revision'],reason='Reviewed',confirmed=True,request_token=str(uuid.uuid4()))
    assert client.post('/api/receipts.gilet.cancel',json=payload,headers={'Origin':ORIGIN}).status_code==403
    assert client.post('/api/receipts.gilet.cancel',json=payload,headers={'Origin':'https://other.test','X-CSRF-Token':session['csrf']}).status_code==403
    A.save_user(employee_name='Receipt canceller',name='mixed_no_accounts',password=PASSWORD,permissions=['payments.view','payments.delete'])
    limited=login(client,'mixed_no_accounts')
    assert post(client,limited,'receipts.gilet.cancel',payload).status_code==403
    session=login(client);before=F.projection(conn,sid);old_audit=F._audit
    def fail(c,student,receipt,action,*args,**kwargs):
        if action=='gilet_item_cancelled':raise ValueError('Injected failure after retained allocation copy')
        return old_audit(c,student,receipt,action,*args,**kwargs)
    monkeypatch.setattr(F,'_audit',fail)
    assert post(client,session,'receipts.gilet.cancel',payload).status_code==422
    assert F.projection(conn,sid)==before
    with conn.cursor() as c:
        c.execute('SELECT COUNT(*) n FROM monthly_gilet_cancellations');assert c.fetchone()['n']==0
        c.execute('SELECT COUNT(*) n FROM monthly_receipts WHERE student_id=%s',(sid,));assert c.fetchone()['n']==1


def test_delete_authorized_worker_can_cancel_without_creating_cash_rights(browser):
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    A.save_user(employee_name='Corrections clerk',name='mixed_clerk',password=PASSWORD,
        permissions=['payments.view','payments.delete','accounts.delete'])
    worker=login(client,'mixed_clerk');result,_=execute(client,worker,rid);assert result.status_code==200,result.text
    assert F.projection(conn,sid)['cash_paid']==4000


@pytest.mark.parametrize('same_token',[False,True])
def test_concurrent_partial_cancellation_is_one_financial_correction(browser,same_token):
    import receipt_item_cancellation as C
    import web_runtime as W
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    p=plan(client,session,rid);auth=A.login('owner',PASSWORD);key=str(uuid.uuid4())
    def run(token):
        with W.request_context(conn.dsn,{'auth_session':auth}):
            try:return C.cancel(rid,actor='owner',reason='Concurrent correction',confirmed=True,token=token,expected_revision=p['revision'])
            except ValueError:return None
    with ThreadPoolExecutor(max_workers=2) as pool:result=list(pool.map(run,[key,key if same_token else str(uuid.uuid4())]))
    assert sum(bool(r and r['created']) for r in result)==1
    assert sum(r is not None for r in result)==(2 if same_token else 1)
    assert F.projection(conn,sid)['cash_paid']==4000


def test_stale_confirmation_and_unsupported_receipt_do_not_write(browser):
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    p=plan(client,session,rid)
    F.take_payment(sid,100,actor='owner')
    response=post(client,session,'receipts.gilet.cancel',dict(id=rid,revision=p['revision'],reason='Stale',confirmed=True,request_token=str(uuid.uuid4())))
    assert response.status_code==422
    F.cancel_receipt(rid,'Whole receipt already voided',actor='owner')
    assert not plan(client,session,rid)['eligible']


@pytest.mark.parametrize('distributed',[False,True])
def test_partial_or_distributed_legacy_garment_is_explicitly_blocked(browser,distributed):
    import db
    client,conn=browser;session=login(client)
    response=post(client,session,'students.create',intake(mode='pay',registration=False,gilet=True,amount='0'))
    assert response.status_code==200,response.text
    sid=response.json()['id']
    # Import-style synthetic data: a genuinely partial garment allocation.
    with conn.cursor() as c:
        A.require(c,'payments.create')
        c.execute('SELECT fee_id FROM monthly_fees WHERE student_id=%s',(sid,));fee=c.fetchone()['fee_id']
        c.execute('SELECT period_id FROM monthly_periods WHERE student_id=%s ORDER BY period_no',(sid,));period=c.fetchone()['period_id']
        c.execute('''INSERT INTO monthly_receipts(student_id,amount,received_on,payer,request_token,request_hash,created_by)
            VALUES(%s,1800,%s,'Synthetic',%s,%s,'owner') RETURNING receipt_id''',(sid,F.today(),str(uuid.uuid4()),'x'*64))
        rid=c.fetchone()['receipt_id']
        c.execute('INSERT INTO monthly_allocations(receipt_id,fee_id,amount) VALUES(%s,%s,1000)',(rid,fee))
        c.execute('INSERT INTO monthly_allocations(receipt_id,period_id,amount) VALUES(%s,%s,800)',(rid,period))
        if distributed:
            c.execute('''INSERT INTO monthly_receipts(student_id,amount,received_on,payer,request_token,request_hash,created_by)
                VALUES(%s,1000,%s,'Synthetic',%s,%s,'owner') RETURNING receipt_id''',(sid,F.today(),str(uuid.uuid4()),'y'*64))
            other=c.fetchone()['receipt_id'];c.execute('INSERT INTO monthly_allocations(receipt_id,fee_id,amount) VALUES(%s,%s,1000)',(other,fee))
    p=plan(client,session,rid)
    assert not p['eligible'] and 'جزئي' in p['reason']
    before=F.projection(conn,sid)
    denied=post(client,session,'receipts.gilet.cancel',dict(id=rid,revision=F.receipt_revision(conn,sid),reason='Unsupported',confirmed=True,request_token=str(uuid.uuid4())))
    assert denied.status_code==422 and F.projection(conn,sid)==before
    state=client.get('/api/snapshot').json();assert not next(r for r in state['receipts'] if r['id']==rid)['gilet_item_eligible']


def test_no_garment_and_standalone_receipts_offer_no_partial_action(browser):
    from tests.test_gilet_exemption import sale
    client,conn=browser;session=login(client)
    response=post(client,session,'students.create',intake(mode='pay',registration=True,gilet=False,amount='4000'))
    sid=response.json()['id'];rid=F.projection(conn,sid)['receipts'][0]['receipt_id']
    assert not plan(client,session,rid)['eligible']
    standalone,_=sale(client,session,sid);assert standalone.status_code==200
    assert not plan(client,session,standalone.json()['id'])['eligible']


def test_cancellation_historical_totals_and_return_to_paying_do_not_restore_obligation(browser):
    import account_records as AR
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    # Existing historical date editor is intentional and keeps the source tariff.
    day=F.today()-timedelta(days=5)
    corrected,_=F.correct_receipt(rid,6000,'Actual past cash date',received_on=day,actor='owner')
    result,_=execute(client,session,corrected);assert result.status_code==200,result.text
    after=F.projection(conn,sid)
    assert after['fees']['gilet']['cancelled']==2000 and after['fees']['gilet']['paid']==0
    # Corrected-ledger reports exclude both cancelled cash and its cancelled charge,
    # even for a payment-date range before the later clerical decision.
    with conn.cursor() as c:
        c.execute('SELECT start_date FROM monthly_accounts WHERE student_id=%s',(sid,));start=c.fetchone()['start_date']
    AR.change_account(sid,actor='owner',reason='Earlier actual enrollment',start_date=day-timedelta(days=10),status='regular')
    historic=F.projection(conn,sid,as_of=day)
    assert historic['fees']['gilet']['cancelled']==2000 and historic['due_now']==0
    report=post(client,session,'reports.finance',dict(start=str(day),end=str(day))).json()
    assert report['summary']['income']==4000 and report['summary']['gilet']==0
    p=post(client,session,'exemption.preview',dict(student=sid)).json()
    exempt=post(client,session,'exemption.apply',dict(student=sid,revision=p['revision'],ledger_revision=p['ledger_revision'],due=p['due'],reason='School exemption',confirmed=True,request_token=str(uuid.uuid4())))
    assert exempt.status_code==200,exempt.text
    AR.change_account(sid,actor='owner',reason='Return to paid',start_date=day-timedelta(days=10),status='regular')
    assert F.projection(conn,sid)['fees']['gilet']['cancelled']==2000


def test_upgrade_17_to_18_preserves_existing_cash_frozen_receipts_photos_and_access(browser):
    import base64,db,schema_migrations as M
    from tests.test_student_photos import picture
    client,conn=browser;session=login(client)
    response=post(client,session,'students.create',intake(mode='pay',registration=True,gilet=True,amount='6000',photo=base64.b64encode(picture()).decode()))
    assert response.status_code==200,response.text
    A.save_user(employee_name='Existing synthetic cashier',name='old_cashier',password=PASSWORD,permissions=['payments.view','payments.create'])
    tables=('students','parents','student_photos','monthly_accounts','monthly_fees','monthly_periods',
        'monthly_receipts','monthly_allocations','monthly_allocation_snapshots','gilet_sales','gilet_deliveries',
        'monthly_waivers','app_users','app_sessions','record_changes')
    before={}
    with conn.cursor() as c:
        c.execute('DROP TABLE monthly_gilet_cancellations')
        c.execute('ALTER TABLE monthly_allocations DROP COLUMN copied_from_id')
        c.execute('DELETE FROM app_schema_migrations WHERE version>17')
        c.execute('INSERT INTO app_schema_migrations(version) VALUES(17) ON CONFLICT DO NOTHING')
        for table in tables:
            c.execute('SELECT * FROM '+table);before[table]=[dict(x) for x in c.fetchall()]
    transaction=db.new_transaction_connection()
    try:assert M.upgrade(transaction)
    finally:transaction.close()
    with conn.cursor() as c:
        for table in tables:
            c.execute('SELECT * FROM '+table);after=[dict(x) for x in c.fetchall()]
            if table=='monthly_allocations':after=[{k:v for k,v in r.items() if k!='copied_from_id'} for r in after]
            assert after==before[table],table
        c.execute('SELECT COUNT(*) n FROM monthly_gilet_cancellations');assert c.fetchone()['n']==0
        c.execute('SELECT MAX(version) v FROM app_schema_migrations');assert c.fetchone()['v']==18


@pytest.mark.parametrize('state',['exempt','archived'])
def test_historical_item_correction_does_not_reactivate_or_change_student_status(browser,state):
    import student_debt as D
    from tests.test_gilet_exemption import exemption
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    if state=='exempt':
        r,_=exemption(client,session,sid);assert r.status_code==200,r.text
    else:D.archive_student(sid,actor='owner')
    before=F.projection(conn,sid)
    result,_=execute(client,session,rid);assert result.status_code==200,result.text
    after=F.projection(conn,sid)
    assert after['account']==before['account'] and after['due_now']==0
    assert after['cash_paid']==4000
    with pytest.raises(ValueError):D.pay_unpaid_fee(sid,'gilet',actor='owner',expected_due=2000)


def test_current_corrected_receipt_keeps_multiple_month_allocations(browser):
    client,conn=browser;session=login(client);sid,rid=create(client,session)
    current,_=F.correct_receipt(rid,7000,'Correct the original amount',actor='owner')
    source=allocations(conn,current);before=F.projection(conn,sid)
    assert len([x for x in source if x['period_id']])==2
    result,_=execute(client,session,current);assert result.status_code==200,result.text
    carried=allocations(conn,result.json()['id'])
    assert [(x['period_id'],x['amount']) for x in carried if x['period_id']]==[(x['period_id'],x['amount']) for x in source if x['period_id']]
    assert F.projection(conn,sid)['periods']==before['periods']


def test_multiple_allocations_for_one_fully_paid_garment_in_same_receipt_are_supported(browser):
    client,conn=browser;session=login(client)
    response=post(client,session,'students.create',intake(mode='pay',registration=False,gilet=True,amount='0'))
    sid=response.json()['id']
    with conn.cursor() as c:
        A.require(c,'payments.create')
        c.execute('SELECT fee_id FROM monthly_fees WHERE student_id=%s',(sid,));fee=c.fetchone()['fee_id']
        c.execute('SELECT period_id FROM monthly_periods WHERE student_id=%s ORDER BY period_no',(sid,));period=c.fetchone()['period_id']
        c.execute('''INSERT INTO monthly_receipts(student_id,amount,received_on,payer,request_token,request_hash,created_by)
            VALUES(%s,5000,%s,'Synthetic',%s,%s,'owner') RETURNING receipt_id''',(sid,F.today(),str(uuid.uuid4()),'z'*64))
        rid=c.fetchone()['receipt_id']
        for value in (800,1200):c.execute('INSERT INTO monthly_allocations(receipt_id,fee_id,amount) VALUES(%s,%s,%s)',(rid,fee,value))
        c.execute('INSERT INTO monthly_allocations(receipt_id,period_id,amount) VALUES(%s,%s,3000)',(rid,period))
    assert plan(client,session,rid)['eligible']
    result,_=execute(client,session,rid);assert result.status_code==200,result.text
    assert F.projection(conn,sid)['cash_paid']==3000 and F.projection(conn,sid)['due_now']==0
