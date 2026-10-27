"""Real PostgreSQL: cash dates, exact drill-down, scopes and verifiable exports."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import uuid

from openpyxl import load_workbook
import pytest

import access_control as A
import expense_finance as E
import monthly_finance as F
import payroll_finance as P
from tests.test_access_control import access, PASSWORD
from tests.test_live_pages_http import browser, intake, login, post


def report(client, session, start=None, end=None):
    today = F.today()
    r = post(client, session, 'reports.finance', dict(start=str(start or today.replace(day=1)), end=str(end or today)))
    assert r.status_code == 200, r.text
    return r.json()


def setup_cash(browser):
    client, conn = browser
    session = login(client)
    today = F.today();start = today.replace(day=1);previous = start - timedelta(days=1)
    student = post(client, session, 'students.create', intake(mode='pay', registration=True, gilet=True, amount='6000')).json()['id']
    F.take_payment(student, 500, actor='owner', as_of=previous)
    expense, _ = E.create(category='=1+1 (supplies)', amount_total=1200, amount_paid=1000,
                          expense_date=previous, payee='Supplier QA')
    E.settle(expense, 200, on_date=today, actor='owner', reason='Later settlement')
    E.create(category='Unpaid QA', amount_total=900, amount_paid=0, expense_date=today)
    teacher = post(client, session, 'tools.change', dict(table='teachers', mode='create', confirmed=True,
        reason='QA', request_token=str(uuid.uuid4()), values=dict(full_name='Private teacher QA', national_id='REPORT-QA',
        salary='3000', hire_date='2020-01-01', department_id=1))).json()['id']
    advance, _ = P.save_advance(teacher, str(today)[:7], 1000, actor='owner', on_date=today)
    salary, _ = P.save_salary(teacher, str(today)[:7], 3000, deductions=1000, actor='owner', on_date=today)
    return client, conn, session, start, previous, student, expense, salary, advance


def test_period_rows_reconcile_cash_and_prior_invoice_settlement(browser):
    client, conn, session, start, previous, student, expense, salary, advance = setup_cash(browser)
    result = report(client, session)
    s = result['summary'];items = result['items']
    assert s['income'] == 6000
    assert s['operational'] == 200  # Cash paid now, despite last month's invoice.
    assert s['salary'] == 2000 and s['advance'] == 1000
    assert s['outgoing'] == 3200 and s['net'] == 2800
    assert s['opening'] == -500 and s['closing'] == 2300
    assert s['registration'] + s['gilet'] + s['monthly'] + s['other_income'] == s['income']
    assert sum(r['amount'] for r in items if r['category'] == 'income') == s['income']
    assert sum(r['amount'] for r in items if r['category'] != 'income') == s['outgoing']
    assert result['debt']['suppliers_total'] == 900
    assert sum(r['amount'] for r in result['debt']['suppliers']) == 900
    legacy = F.financial_report(conn, start, F.today())
    assert s['income'] == legacy['total_cash']
    assert s['outgoing'] == legacy['expenses']
    assert result['debt']['students_total'] == legacy['due_now']
    before = report(client, session, previous, previous)
    assert before['summary']['operational'] == 1000
    assert before['debt']['suppliers_total'] == 200
    assert s['prepaid'] <= s['monthly']


def test_corrections_and_document_void_do_not_create_fake_cash(browser):
    client, conn, session, start, previous, student, expense, salary, advance = setup_cash(browser)
    P.amend('teacher_payments', salary, actor='owner', reason='Void document only', cancel=True)
    result = report(client, session)
    assert result['summary']['salary'] == 2000
    assert next(r for r in result['items'] if r['category'] == 'salary')['void_document']
    E.cancel(expense, actor='owner', reason='Erroneous expense')
    result = report(client, session)
    assert result['summary']['operational'] == 0
    assert not any(r['category'] == 'operational' for r in result['items'])
    with conn.cursor() as c:
        c.execute("SELECT receipt_id FROM monthly_receipts WHERE student_id=%s AND received_on=%s AND state='active'", (student, F.today()))
        receipt = c.fetchone()['receipt_id']
    F.cancel_receipt(receipt, 'Wrong entry', actor='owner')
    result = report(client, session)
    assert result['summary']['income'] == 0
    assert result['summary']['closing'] == result['summary']['opening'] - result['summary']['outgoing']


def test_export_matches_display_filters_safe_text_and_stale_revision(browser):
    client, conn, session, *_ = setup_cash(browser)
    current = report(client, session)
    request = dict(start=current['start'], end=current['end'], revision=current['revision'], category='operational', query='(')
    downloaded = post(client, session, 'reports.finance.export', request)
    assert downloaded.status_code == 200, downloaded.text[:200] if downloaded.status_code != 200 else ''
    book = load_workbook(BytesIO(downloaded.content))
    assert book.sheetnames == ['الملخص', 'الحركات', 'الديون', 'الطلاب الجدد', 'حضور الطلاب', 'مقارنة الفترتين']
    rows = list(book['الحركات'].values)
    assert len(rows) == 2 and rows[1][2] == '=1+1 (supplies)' and rows[1][5] == 200
    assert book['الحركات']['C2'].data_type == 's'
    summary = {r[0]: r[1] for r in book['الملخص'].values if len(r) >= 2}
    assert summary['المقبوض خلال الفترة'] == current['summary']['income']
    assert summary['صافي الحركات المعروضة'] == -200
    E.create(category='New movement', amount_total=50, amount_paid=50, expense_date=F.today())
    stale = post(client, session, 'reports.finance.export', request)
    assert stale.status_code == 422 and 'حدّث' in stale.json()['error']


def test_report_permission_and_detail_redaction_are_enforced_on_server(browser):
    client, conn, session, *_ = setup_cash(browser)
    def account(username, permissions):
        r = post(client, session, 'users.save', dict(name=username, username=username, password=PASSWORD,
                                                   permissions=permissions, active=True))
        assert r.status_code == 200, r.text
    account('report_only', ['reports.view'])
    account('no_reports', ['students.view'])
    scoped = login(client, 'report_only')
    result = report(client, scoped)
    assert result['summary']['outgoing'] == 3200
    assert all(not r['label'] and not r['party'] and r['source_id'] is None and r['student_id'] is None for r in result['items'])
    assert result['debt']['students'] == result['debt']['suppliers'] == []
    assert 'Private teacher QA' not in str(result) and 'Supplier QA' not in str(result)
    denied = login(client, 'no_reports')
    r = post(client, denied, 'reports.finance', dict(start=str(F.today()), end=str(F.today())))
    assert r.status_code == 403


def test_empty_range_invalid_dates_and_inclusive_boundaries(browser):
    client, conn = browser;session = login(client)
    empty = report(client, session)
    assert empty['items'] == [] and empty['summary']['income'] == empty['summary']['outgoing'] == 0
    E.create(category='Boundary', amount_total='12.50', amount_paid='12.50', expense_date=F.today())
    result = report(client, session, F.today(), F.today())
    assert Decimal(str(result['summary']['outgoing'])) == Decimal('12.50')
    for data in (dict(start='2025-02-29', end=str(F.today())), dict(start=str(F.today()), end='2099-01-01'),
                 dict(start=str(F.today()), end='2020-01-01')):
        assert post(client, session, 'reports.finance', data).status_code == 422


def test_concurrent_change_cannot_mix_totals_and_debt_snapshot(browser, monkeypatch):
    client, conn = browser;session = login(client)
    original = F.projections
    injected = False
    def during_read(connection, *args, **kwargs):
        nonlocal injected
        if not injected:
            injected = True
            E.create(category='Concurrent unpaid invoice', amount_total=700, amount_paid=0, expense_date=F.today())
        return original(connection, *args, **kwargs)
    monkeypatch.setattr(F, 'projections', during_read)
    first = report(client, session)
    assert injected and first['summary']['outgoing'] == 0 and first['debt']['suppliers_total'] == 0
    second = report(client, session)
    assert second['debt']['suppliers_total'] == 700


def test_secondary_reports_require_only_relevant_filters(browser):
    client, conn = browser;session = login(client)
    roster = post(client, session, 'reports.read', dict(kind='students'))
    assert roster.status_code == 200, roster.text
    debt = post(client, session, 'reports.read', dict(kind='debt', end=str(F.today())))
    assert debt.status_code == 200, debt.text
