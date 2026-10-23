"""French school workflow through the production API on disposable PostgreSQL/TLS."""
import io
import os
import json
from pathlib import Path
import sys
from pypdf import PdfReader
from openpyxl import load_workbook
from playwright.async_api import async_playwright, expect
from tests.date_entry_helpers import fill_field


async def main(origin,username,password):
    out=Path(__file__).resolve().parents[1]/'visual-evidence'/'french-live'
    out.mkdir(parents=True,exist_ok=True)
    errors=[]
    async with async_playwright() as pw:
        browser=await pw.chromium.launch(headless=True,channel='chrome' if sys.platform=='win32' else None)
        context=await browser.new_context(ignore_https_errors=True,viewport={'width':1440,'height':1000},accept_downloads=True)
        page=await context.new_page()
        page.on('pageerror',lambda e:errors.append(str(e)))
        try:
            await page.goto(origin,wait_until='networkidle')
            await page.locator('[data-language="fr"]').click()
            await page.wait_for_load_state('networkidle')
            await page.locator('#live-login [name="username"]').fill(username)
            await page.locator('#live-login [name="password"]').fill('wrong-password')
            await page.locator('#preview-enter').click()
            await expect(page.locator('#live-login .error')).to_contain_text('Connexion impossible')
            await page.locator('#live-login [name="password"]').fill(password)
            await page.locator('#preview-enter').click()
            await expect(page.locator('#preview-login')).to_be_hidden()
            await expect(page.locator('[data-language="ar"]')).to_be_hidden()
            await page.locator('nav [data-page="registration"]').click()
            await page.locator('#department-form [name="name"]').fill('Les Étoiles — النجوم')
            await page.locator('#department-form button').click()
            await expect(page.locator('#toast')).to_contain_text('Opération enregistrée')
            data=await page.evaluate('liveCall("snapshot")')
            section=next(d['id'] for d in data['departments'] if d['name']=='Les Étoiles — النجوم')
            await page.locator('nav [data-page="students"]').click()
            await page.locator('#add-student').click()
            for key,value in {'name':'Élève مريم','birth':'2022-01-01','mother':'Mère فاطمة','motherPhone':'22334455','amount':'4500'}.items():
                await fill_field(page.locator('#student-form [name="'+key+'"]'), value)
            await page.locator('#student-form [name="gender"]').select_option('أنثى')
            await page.locator('#student-form [name="department_id"]').select_option(str(section))
            await page.locator('#student-form .form-end button.primary').click()
            await expect(page.locator('#student-form')).to_have_count(0)
            await page.locator('#student-search').fill('Élève مريم')
            data=await page.evaluate('liveCall("snapshot")')
            child=next(s for s in data['students'] if s['name']=='Élève مريم')
            assert child['gender']=='أنثى' and child['account']['paid']==4500
            await page.locator('[data-profile="'+child['id']+'"]').click()
            await expect(page.locator('.file-tags')).to_contain_text('Fille')
            await page.locator('#open-profile-edit').click()
            await page.locator('#profile-edit-child [name="health"]').fill('Asthme — حساسية')
            await page.locator('#profile-edit-child button.primary').click()
            await expect(page.locator('#profile-edit-dialog')).to_be_hidden()
            assert next(s for s in (await page.evaluate('liveCall("snapshot")'))['students'] if s['id']==child['id'])['health']=='Asthme — حساسية'
            await page.locator('nav [data-page="payments"]').click()
            await page.locator('[data-finance="receipts"]').click()
            rid=next(r['id'] for r in (await page.evaluate('liveCall("snapshot")'))['receipts'] if r['student']==child['id'])
            await page.locator('#receipt-search').fill('Élève مريم')
            await page.locator('[data-receipt="'+str(rid)+'"]').click()
            await expect(page.locator('#receipt-page')).to_contain_text('Élève مريم')
            async with page.expect_download() as event:
                await page.locator('#receipt-page a').click()
            receipt=await event.value
            await receipt.save_as(out/'receipt-fr.pdf')
            pdf='\n'.join(p.extract_text() for p in PdfReader(out/'receipt-fr.pdf').pages)
            assert 'Reçu de paiement' in pdf and 'MRO' in pdf and 'Élève' in pdf
            await page.screenshot(path=str(out/'receipt-fr.png'),full_page=True)
            await page.locator('nav [data-page="expenses"]').click()
            for key,value in {'category':'Cahiers — دفاتر','amount':'1200','paid':'1000'}.items():
                if key=='paid':await page.locator('#expense-form [name="status"]').select_option('partial')
                await page.locator('#expense-form [name="'+key+'"]').fill(value)
            await page.locator('#expense-form button.primary').click()
            await expect(page.locator('table')).to_contain_text('Cahiers — دفاتر')
            await page.locator('[data-record-type="expenses"][data-record-mode="edit"]').first.click()
            await page.locator('#record-edit-form [name="amount"]').fill('800')
            await page.locator('#record-edit-form [name="reason"]').fill('Correction')
            await page.locator('#record-edit-form [name="agree"]').check()
            await page.locator('#record-edit-form button.primary').click()
            await expect(page.locator('#record-error')).to_contain_text('dépasse le total')
            await expect(page.locator('#record-edit-form [name="paid"]')).to_have_value('1000')
            await page.locator('#record-edit-form [name="amount"]').fill('1400')
            await page.locator('#record-edit-form button.primary').click()
            await expect(page.locator('#record-dialog')).to_be_hidden()
            await page.locator('#nav-teachers').click()
            await page.locator('#add-teacher').click()
            for key,value in {'full_name':'Enseignant Noël','national_id':'FR-TE-101','salary':'2000'}.items():
                await page.locator('#teacher-form [name="'+key+'"]').fill(value)
            await page.locator('#teacher-section').select_option(str(section))
            await page.locator('#teacher-form [name="reason"]').fill('Création pour vérification')
            await page.locator('#teacher-form [name="confirmed"]').check()
            await page.locator('#teacher-form button.primary').click()
            await expect(page.locator('#teacher-roster')).to_contain_text('Enseignant Noël')
            await page.locator('#nav-workspace').click()
            await page.locator('[data-tool-tab="payroll"]').click()
            await page.locator('#new-salary').click()
            await page.locator('#payroll-form [name="teacher"]').select_option(label='Enseignant Noël · FR-TE-101')
            await page.locator('#payroll-form [name="bonus"]').fill('500')
            await page.locator('#payroll-form [name="deductions"]').fill('100')
            await expect(page.locator('#salary-net')).to_contain_text('2\u202f400')
            await page.locator('#payroll-form [name="reason"]').fill('Versement du salaire')
            await page.locator('#payroll-form [name="confirmed"]').check()
            await page.locator('#payroll-form button.primary').click()
            await expect(page.locator('#record-dialog')).to_be_hidden()
            await page.locator('[data-tool-tab="reports"]').click()
            await page.locator('#report-form [name="kind"]').select_option('students')
            await page.locator('#report-form button').click()
            await expect(page.locator('#report-output th').first).to_have_text('Élève')
            async with page.expect_download() as event:await page.locator('#report-xlsx').click()
            download=await event.value
            await download.save_as(out/'report-fr.xlsx')
            workbook=load_workbook(out/'report-fr.xlsx')
            assert workbook.sheetnames==['Rapport'] and workbook.active['A1'].value=='Élève'
            await page.locator('nav [data-page="admin"]').click()
            legends=await page.locator('.grant-row legend').all_text_contents()
            assert not any(any('\u0600'<=c<='\u06ff' for c in text) for text in legends),legends
            await page.locator('#demo-user-create [name="name"]').fill('Agent français')
            await page.locator('#demo-user-create [name="username"]').fill('agent_fr')
            await page.locator('#demo-user-create [name="password"]').fill('0000')
            await page.locator('#demo-user-create [value="student_attendance.create"]').check()
            await page.locator('#demo-user-create [value="departments.view"]').check()
            await page.locator('#demo-user-create button.primary').click()
            await expect(page.locator('#toast')).to_contain_text('enregistrée')
            await page.locator('#demo-logout').click()
            await expect(page.locator('[data-language="fr"]')).to_be_visible()
            await page.locator('#live-login [name="username"]').fill('agent_fr')
            await page.locator('#live-login [name="password"]').fill('0000')
            await page.locator('#preview-enter').click()
            await expect(page.locator('#preview-login')).to_be_hidden()
            assert await page.evaluate('liveData.receipts.length===0&&liveData.overview===null&&liveData.students.every(s=>!s.account&&!s.birth&&!s.health)')
            await page.locator('#nav-student-attendance').click()
            await page.locator('#attendance-department').select_option(str(section))
            await page.locator('#start-attendance').click()
            await expect(page.locator('[data-roll-row]')).to_have_count(1)
            await expect(page.locator('[data-roll-row]').first).to_contain_text('Élève مريم')
            await expect(page.locator('[data-roll-status="present"]').first).to_be_enabled()
            await page.locator('[data-roll-status="present"]').first.click()
            await page.locator('#roll-review').click()
            await page.locator('#roll-save').click()
            await expect(page.locator('#toast')).to_contain_text('validée et enregistrée')
            await page.reload(wait_until='networkidle')
            await expect(page.locator('html')).to_have_attribute('dir','ltr')
            await page.locator('#demo-logout').click()
            await page.locator('[data-language="ar"]').click()
            await page.wait_for_load_state('networkidle')
            await expect(page.locator('html')).to_have_attribute('dir','rtl')
            assert not errors,errors
            print('FRENCH_LIVE_OK login intake profile payment receipt PDF expense validation XLSX roles attendance logout language isolation')
        except Exception:
            await page.screenshot(path=str(out/'failure.png'),full_page=True)
            (out/'failure.txt').write_text((await page.locator('body').inner_text())+'\n'+json.dumps(errors),encoding='utf-8')
            raise
        finally:
            await browser.close()


if __name__=='__main__':
    import asyncio
    state=json.loads((Path(os.environ['KG_LIVE_QA_DIR'])/'state.json').read_text())
    from tests.qa_live_pages import guard
    guard(state['dsn'])
    asyncio.run(main(state['origin'],state['username'],state['password']))
