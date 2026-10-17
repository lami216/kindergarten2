"""Reproduce the independent review's debt, lost-update, export and identity cases."""
import io
import json
import uuid
import zipfile

import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest
from openpyxl import load_workbook

import db
import monthly_finance as F
import records as R
import student_debt as D
from account_records import change_account, change_waiver
from data_export import build_export, verify_export
from student_search import child_labels
from tests.test_record_management import ledger, child, row, scalar
from tests.test_family_ui import SCRIPT


def test_reopened_debt_on_exempt_child_visible_and_blocks_both_archive_paths(ledger):
    conn, marker = ledger
    student = child(conn, marker)
    D.waive_debt(student, 'قرار أول', True, actor='review')
    change_account(student, start_date=F.today(), status='exempt', actor='review', reason='إعفاء لاحق')
    ident = scalar(conn, "SELECT w.waiver_id FROM monthly_waivers w JOIN monthly_fees f ON f.fee_id=w.fee_id WHERE w.student_id=%s AND f.kind='registration'", (student,))
    baseline = F.financial_overview(conn)['due_now']
    change_waiver(ident, cancel=True, actor='review', reason='إسقاط خاطئ')
    assert F.projection(conn, student)['due_now'] == 1000
    assert F.financial_overview(conn)['due_now'] == baseline + 1000
    for archive in (lambda: D.archive_student(student, actor='review'),
                    lambda: R.cancel('students', student, actor='review', reason='أرشفة'),
                    lambda: change_account(student, cancel=True, actor='review', reason='إلغاء')):
        with pytest.raises(ValueError):
            archive()
    assert row(conn, 'students', student)['archived_at'] is None
    assert D.waive_debt(student, 'قرار صريح جديد', True, actor='review')[0] == 1000
    D.archive_student(student, actor='review')
    assert F.projection(conn, student)['due_now'] == 0


def test_exempt_blocks_new_cash_but_allows_historical_receipt_void_and_replay(ledger):
    conn, marker = ledger
    student = child(conn, marker, cash=4000)
    receipt = scalar(conn, 'SELECT receipt_id FROM monthly_receipts WHERE student_id=%s', (student,))
    change_account(student, start_date=F.today(), status='exempt', actor='review', reason='إعفاء لاحق')
    with pytest.raises(ValueError, match='معفي'):
        F.take_payment(student, 1000, actor='review')
    count = scalar(conn, 'SELECT COUNT(*) FROM monthly_periods WHERE student_id=%s', (student,))
    token = str(uuid.uuid4())
    assert F.cancel_receipt(receipt, 'قبض أدخل خطأ', actor='review', token=token)
    assert not F.cancel_receipt(receipt, 'قبض أدخل خطأ', actor='review', token=token)
    assert F.projection(conn, student)['due_now'] == 4000
    assert F.projection(conn, student)['cash_paid'] == 0
    assert scalar(conn, 'SELECT COUNT(*) FROM monthly_periods WHERE student_id=%s', (student,)) == count


def test_old_child_form_cannot_erase_new_medical_information(ledger):
    conn, marker = ledger
    student = child(conn, marker)
    app = AppTest.from_string(SCRIPT).run(timeout=30)
    name = row(conn, 'students', student)['full_name']
    app.selectbox(key='student_select_edit').set_value((student, name)).run(timeout=30)
    app.button(key='open_dialog_student_'+student).click().run(timeout=30)
    R.edit('students', student, {'medical_details': 'حساسية جديدة'}, actor='other session', reason='معلومة جديدة')
    app.text_input(key=f'edit_{student}_address').set_value('عنوان جديد')
    next(b for b in app.button if b.label == '💾 حفظ التعديلات').click().run(timeout=30)
    assert not app.exception
    assert any('تغيّر' in e.value for e in app.error)
    assert row(conn, 'students', student)['medical_details'] == 'حساسية جديدة'
    assert row(conn, 'students', student)['address'] is None
    next(b for b in app.button if b.label=='إغلاق دون حفظ').click().run(timeout=30)
    next(b for b in app.button if b.label=='تجاهل التغييرات والخروج').click().run(timeout=30)
    app.button(key='open_dialog_student_'+student).click().run(timeout=30)
    assert app.text_area(key=f'edit_{student}_health').value == 'حساسية جديدة'
    app.text_input(key=f'edit_{student}_address').set_value('عنوان جديد')
    next(b for b in app.button if b.label == '💾 حفظ التعديلات').click().run(timeout=30)
    assert not app.exception
    saved = row(conn, 'students', student)
    assert saved['medical_details'] == 'حساسية جديدة' and saved['address'] == 'عنوان جديد'


def test_complete_export_includes_monthly_cash_and_full_json_history(ledger):
    conn, marker = ledger
    student = child(conn, marker, cash=12000)
    long_note = ('نص ' * 12000).strip()
    R.edit('students', student, {'general_notes': long_note}, actor='export QA', reason='نص طويل')
    payload, manifest = build_export()
    assert verify_export(payload) == manifest
    with conn.cursor() as c:
        c.execute("SELECT table_name FROM information_schema.tables WHERE table_schema=current_schema() AND table_type='BASE TABLE'")
        assert {r['table_name'] for r in c.fetchall()} == {t['name'] for t in manifest['tables']}
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        receipts = json.loads(archive.read('data/monthly_receipts.json'))
        assert sum(r['amount'] for r in receipts if r['student_id'] == student) == 12000
        history = json.loads(archive.read('data/record_changes.json'))
        assert any(r['record_key'] == student and r['after_state'].get('general_notes') == long_note for r in history if r['after_state'])
        excel = load_workbook(io.BytesIO(archive.read('records.xlsx')))
        sheet = excel['الدفعات الشهرية']
        headers = [c.value for c in sheet[1]]
        assert any(r[headers.index('student_id')] == student and r[headers.index('amount')] == 12000 for r in sheet.iter_rows(min_row=2, values_only=True))
        broken = io.BytesIO()
        with zipfile.ZipFile(broken, 'w') as damaged:
            for name in archive.namelist():
                damaged.writestr(name, b'[]' if name == 'data/monthly_receipts.json' else archive.read(name))
        with pytest.raises(ValueError):
            verify_export(broken.getvalue())


def test_export_failure_is_reported_instead_of_skipping_a_table(ledger, monkeypatch):
    conn, marker = ledger
    child(conn, marker)
    def fail(value):
        raise RuntimeError('export failure')
    monkeypatch.setattr('data_export._cell', fail)
    with pytest.raises(RuntimeError, match='export failure'):
        build_export()


@pytest.mark.parametrize('table', ['teacher_payments', 'teacher_advances'])
def test_payroll_document_and_cash_voids_are_distinct_and_stale_forms_fail(ledger, table):
    import payroll_finance as P
    conn, marker = ledger
    teacher = 'T'+marker
    with conn.cursor() as c:
        c.execute('INSERT INTO teachers(national_id,full_name,salary) VALUES(%s,%s,1000)', (teacher, marker))
    create = P.save_salary if table == 'teacher_payments' else P.save_advance
    ident, _ = create(teacher, str(F.today())[:7], 1000, actor='review')
    previous = row(conn, table, ident)
    P.amend(table, ident, month=str(F.today())[:7], notes='معلومة أحدث', actor='other', reason='تعديل آخر')
    with pytest.raises(ValueError, match='تغيّر'):
        P.amend(table, ident, cancel=True, actor='review', reason='نموذج قديم', expected=R.revision(previous))
    fresh = row(conn, table, ident)
    P.amend(table, ident, cancel=True, actor='review', reason='إلغاء المستند فقط', expected=R.revision(fresh))
    assert row(conn, table, ident)['cash_amount'] == 1000
    P.correct_cash_entry(table, ident, cancel=True, actor='review', reason='إلغاء إدخال مالي خاطئ')
    assert row(conn, table, ident)['cash_amount'] == 0


def test_session_connections_and_audit_identity_cannot_interleave(ledger):
    conn, marker = ledger
    script = '''
import streamlit as st, db, record_history as H
from unittest.mock import patch
with patch.object(db.st,'secrets',{'postgres':{'url':__import__('os').environ['TEST_POSTGRES_DSN']}}):
 connection=db.get_connection()
 connection.rollback()
 with connection.cursor() as c:
  c.execute('SELECT pg_backend_pid() AS pid')
  st.session_state['pid']=c.fetchone()['pid']
  H.context(c,st.session_state['auth_user'],local=False)
 if st.button('write'):
  with connection.cursor() as c:
   c.execute('INSERT INTO school_departments(name) VALUES(%s)',(st.session_state['name'],))
'''
    apps=[]
    for actor in ('session A', 'session B'):
        app=AppTest.from_string(script)
        app.session_state['auth_user']=actor
        app.session_state['name']=marker+actor
        app.run(timeout=30)
        assert not app.exception
        apps.append(app)
    assert apps[0].session_state['pid'] != apps[1].session_state['pid']
    # B has just set its context; A must still write as A, even on the next rerun.
    for app in apps:
        app.button[0].click().run(timeout=30)
        assert not app.exception
    with conn.cursor() as c:
        c.execute("SELECT after_state->>'name' AS name,actor FROM record_changes WHERE table_name='school_departments' AND after_state->>'name' IN (%s,%s)", (marker+'session A', marker+'session B'))
        assert {r['name']:r['actor'] for r in c.fetchall()} == {marker+'session A':'session A',marker+'session B':'session B'}


def test_duplicate_children_have_human_context_and_last_resort_identity():
    labels=child_labels(pd.DataFrame([
        dict(student_id='A',full_name='أحمد',mother_name='أم أحمد',mother_mobile='1111'),
        dict(student_id='B',full_name='أحمد',father_name='والد أحمد',father_mobile='2222'),
        dict(student_id='C',full_name='محمد'),dict(student_id='D',full_name='محمد')]))
    assert 'أم أحمد' in labels['A'] and '1111' in labels['A']
    assert 'والد أحمد' in labels['B'] and labels['A'] != labels['B']
    assert labels['C'] != labels['D']
