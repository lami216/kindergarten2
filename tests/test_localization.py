from concurrent.futures import ThreadPoolExecutor
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import re

import pytest
import localization as L


def test_catalog_placeholders_and_frontend_keys():
    from scripts.check_translations import validate
    assert validate() == []


def test_every_live_resource_field_and_table_has_french_copy():
    import access_control as A
    import live_tools
    labels = [value[0] for value in A.RESOURCES.values()]
    labels += list(live_tools.LABELS.values())
    labels += [value[0] for value in live_tools.TABLES.values()]
    assert not [label for label in labels if label not in L.BY_SOURCE]
    with L.using_language('fr'):
        assert not [label for label in labels if re.search('[\u0600-\u06ff]',L.label(label))]


def test_language_is_request_local_and_restores_after_error():
    import threading
    barrier = threading.Barrier(2)
    def translate(lang):
        with L.using_language(lang):
            barrier.wait(timeout=5)
            return L.label('الطلاب')
    with ThreadPoolExecutor(2) as pool:
        assert list(pool.map(translate,['fr','ar'])) == ['Élèves','الطلاب']
    with pytest.raises(ValueError), L.using_language('fr'):
        raise ValueError()
    assert L.language() == 'ar'
    with L.using_language('fr'):
        assert L.message('لا يمكن حذف الطالب؛ عليه مبلغ 1,000 أوقية قديمة.') == 'Impossible de supprimer l’élève : il doit 1,000 MRO.'
        assert not re.search('[\u0600-\u06ff]',L.message('خطأ جديد غير مدعوم'))
        assert L.message('لا يمكن حذف ولي الأمر؛ الطفل الطلاب عليه مبلغ 100 أوقية قديمة.') == 'Impossible de supprimer le parent : l’élève الطلاب doit 100 MRO.'


def test_excel_translates_headers_and_enums_without_touching_user_text():
    import pandas as pd
    from openpyxl import load_workbook
    from sections.reports import _to_excel
    data=[{'الطالب':'=1+1','المستلم':'الطلاب','الحالة':'معفي','المبلغ':1200}]
    original=json.dumps(data,ensure_ascii=False)
    with L.using_language('fr'):
        output=_to_excel(pd.DataFrame(L.report_rows(data)),sheet_name='Rapport')
    workbook=load_workbook(BytesIO(output))
    assert workbook.sheetnames == ['Rapport']
    sheet=workbook.active
    assert [c.value for c in sheet[1]] == ['Élève','Bénéficiaire','Statut', 'Montant']
    assert [c.value for c in sheet[2]] == ['=1+1','الطلاب','Exonéré',1200]
    assert sheet['A2'].data_type == 's'
    assert json.dumps(data,ensure_ascii=False) == original


def test_api_login_validation_language_and_catalog_assets():
    from starlette.testclient import TestClient
    from live_pages import make_app
    with TestClient(make_app('unused','https://testserver'),base_url='https://testserver') as client:
        for language,expected in [('fr','Connexion impossible'),('ar','تعذّر الدخول')]:
            result=client.post('/api/login',headers={'Origin':'https://testserver','X-App-Language':language},json={'username':'x','password':''})
            assert result.status_code == 401
            assert expected in result.json()['error']
        client.cookies.set('kg_language','fr')
        assert 'Connectez-vous' in client.get('/api/snapshot').json()['error']
        assert client.get('/locales/messages.js').status_code == 200
        html=client.get('/').text
        assert html.index('i18n.js') < html.index('preview.js')
        assert client.get('/i18n.css').status_code == 200


def test_french_pdf_preserves_original_periods_and_user_content(monkeypatch):
    from copy import deepcopy
    from pypdf import PdfReader
    import receipt_pdf
    record={'receipt':{'receipt_id':42,'state':'cancelled','corrects_id':None,'student_name':'Alice élève — مريم',
        'payer':'Parent Dupont','received_on':date(2026,10,15),'amount':4500,'cancel_reason':'Erreur de saisie',
        'cancelled_by':'owner','cancelled_at':None},'replacements':[],
        'items':[{'fee_kind':'monthly','period_no':0,'starts':date(2026,10,15),'ends':date(2026,11,15),'amount':4500}]}
    original=deepcopy(record)
    monkeypatch.setattr(receipt_pdf,'receipt_detail',lambda *args,**kwargs:deepcopy(record))
    pdf=receipt_pdf.render_receipt_pdf(None,42,locale='fr')
    text='\n'.join(p.extract_text() for p in PdfReader(BytesIO(pdf)).pages)
    for expected in ('Reçu de paiement','Alice','Parent Dupont','2026-10-15','2026-11-14','MRO','Annulé'):
        assert expected in text
    assert '4 500 MRO' in text
    assert '\x00' not in text
    assert record == original
    assert L.language() == 'ar'
    out=Path(__file__).resolve().parents[1]/'visual-evidence'/'french'
    out.mkdir(parents=True,exist_ok=True)
    (out/'receipt-fr.pdf').write_bytes(pdf)
