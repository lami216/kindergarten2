"""Financial report UX, exports and races through the disposable API adapters."""
import asyncio
from datetime import date, timedelta
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import sys
import tempfile
import threading

from openpyxl import load_workbook
from playwright.async_api import async_playwright, expect
from scripts.build_pages_preview import build
from tests.date_entry_helpers import fill_field


async def main():
    out = Path(__file__).resolve().parents[1] / 'visual-evidence' / 'period-reports'
    out.mkdir(parents=True, exist_ok=True)
    errors, passed = [], []
    with tempfile.TemporaryDirectory() as tmp:
        build(Path(tmp))
        server = ThreadingHTTPServer(('127.0.0.1', 0), partial(SimpleHTTPRequestHandler, directory=tmp))
        threading.Thread(target=server.serve_forever, daemon=True).start()
        async with async_playwright() as pw:
            browser = await pw.chromium.launch(headless=True, channel='chrome' if sys.platform == 'win32' else None)
            try:
                for language in ('ar', 'fr'):
                    page = await browser.new_page(viewport={'width': 1440, 'height': 1050})
                    page.on('pageerror', lambda e: errors.append(str(e)))
                    await page.goto(f'http://127.0.0.1:{server.server_port}/', wait_until='networkidle')
                    if language == 'fr': await page.locator('[data-language="fr"]').click()
                    await page.locator('#preview-enter').click()
                    await expect(page.locator('#preview-login')).to_be_hidden()
                    await page.locator('#nav-reports').click()
                    await expect(page.locator('#finance-report-form')).to_be_visible()
                    # An old invoice settled today must contribute only today's 200.
                    await page.evaluate('''async()=>{
                      const first=new Date(today+'T12:00:00');first.setDate(0);
                      const previous=first.toISOString().slice(0,10);
                      const r=await liveCall('expenses.create',{category:'=1+1 (QA)',amount:'1200',paid:'1000',date:previous,payee:'Old supplier',request_token:crypto.randomUUID()});
                      const e=previewState.expenses.find(x=>x.id===r.id);window.settledExpense=r.id;
                      await liveCall('expenses.settle',{id:r.id,revision:e.revision,amount:'200',date:today,reason:'QA',confirmed:true,request_token:crypto.randomUUID()});
                    }''')
                    await page.locator('#finance-report-form button.primary').click()
                    await page.wait_for_function('financeData!==null')
                    assert await page.evaluate('financeData.items.find(r=>r.category==="operational"&&r.source_id===settledExpense).amount') == 200
                    assert await page.evaluate('financeData.summary.income===financeData.items.filter(r=>r.category==="income").reduce((n,r)=>n+r.amount,0)')
                    assert await page.evaluate('financeData.summary.outgoing===financeData.items.filter(r=>r.category!=="income").reduce((n,r)=>n+r.amount,0)')
                    assert await page.evaluate('financeData.summary.closing===financeData.summary.opening+financeData.summary.net')
                    await expect(page.locator('[data-tool-tab="reports"]')).to_have_count(0)
                    await page.screenshot(path=str(out / (language + '-desktop.png')), full_page=True)
                    passed.append(language + ': dedicated route and period details reconcile')

                    await page.locator('[data-report-section=money][role=tab]').click()
                    await page.locator('#fin-query').fill('(')
                    assert await page.locator('.fin-screen-table tbody tr').count() == 1
                    assert await page.evaluate('financeRows()[0].amount') == 200
                    async with page.expect_download() as download_info:
                        await page.locator('#fin-xlsx').click()
                    download = await download_info.value
                    export_path = out / (language + '-filtered.xlsx')
                    await download.save_as(export_path)
                    book = load_workbook(export_path)
                    sheet = book.worksheets[1]
                    assert sheet.max_row == 2 and sheet['C2'].value == '=1+1 (QA)'
                    assert sheet['C2'].data_type == 's' and sheet['F2'].value == 200
                    assert book.worksheets[0].sheet_view.rightToLeft == (language == 'ar')
                    passed.append(language + ': literal search and real safe XLSX match filtered rows')

                    await page.evaluate("liveCall('expenses.create',{category:'New QA',amount:'10',paid:'10',date:today,request_token:crypto.randomUUID()})")
                    await page.locator('#fin-xlsx').click()
                    await expect(page.locator('#toast')).to_have_text(await page.evaluate("RP('exportChanged')"))
                    await page.locator('#finance-report-form button.primary').click()
                    await page.wait_for_function('financeData!==null')
                    await page.locator('#fin-query').fill('')
                    await page.locator('[data-fin-period="lastMonth"]').click()
                    await page.wait_for_function('financeData!==null')
                    assert await page.evaluate('financeData.items.find(r=>r.category==="operational"&&r.source_id===settledExpense).amount') == 1000
                    passed.append(language + ': old invoice cash date and stale-export rejection')

                    # A slower response must never replace a newer selected period.
                    await page.evaluate('''()=>{const original=previewOperation;window.holdNextFinance=true;
                      previewOperation=async(op,data)=>{const value=await original(op,data);
                        if(op==='reports.finance'&&holdNextFinance){holdNextFinance=false;return new Promise(resolve=>window.releaseOldFinance=()=>resolve(value));}return value;};}''')
                    await page.locator('#finance-report-form button.primary').click()
                    await page.wait_for_function('typeof releaseOldFinance==="function"')
                    await page.locator('[data-fin-period="today"]').click()
                    await page.wait_for_function('financeData&&financeData.start===today')
                    await page.evaluate('releaseOldFinance()')
                    await expect(page.locator('#fin-start')).to_have_value(await page.evaluate('today'))
                    assert await page.evaluate('financeData.start===today&&financeData.end===today')
                    await page.locator('[data-fin-period=custom]').click()
                    start = page.locator('#fin-start')
                    await start.locator('xpath=preceding-sibling::*[1]').locator('[data-date-part="year"]').fill('2')
                    assert not await page.locator('#finance-report-form').evaluate('(f)=>f.checkValidity()')
                    await expect(page.locator('#fin-xlsx')).to_have_count(0)
                    await page.locator('[data-fin-period="thisMonth"]').click()
                    await page.wait_for_function('financeData!==null')
                    passed.append(language + ': stale responses and incomplete-date protection')

                    await page.locator('[data-report-mode="other"]').click()
                    await page.locator('#report-form [name="kind"]').select_option('students')
                    await expect(page.locator('#report-form [name="start"]').locator('..')).to_be_hidden()
                    await expect(page.locator('#report-form [name="teacher"]').locator('..')).to_be_hidden()
                    await page.locator('#report-form [name="kind"]').select_option('salaries')
                    await expect(page.locator('#report-form [name="teacher"]').locator('..')).to_be_visible()
                    await page.locator('#report-form [name="start"]').locator('xpath=preceding-sibling::*[1]').locator('[data-date-part="year"]').fill('2')
                    await page.locator('#report-form [name="kind"]').select_option('students')
                    await page.locator('#report-form button').click()
                    await expect(page.locator('#report-output table')).to_be_visible()
                    await page.locator('[data-report-mode="finance"]').click()
                    await page.set_viewport_size({'width': 390, 'height': 844})
                    await page.evaluate("document.querySelector('#sidebar').classList.remove('open')")
                    await page.wait_for_function("(()=>{const r=document.querySelector('#sidebar').getBoundingClientRect();return r.right<=0||r.left>=innerWidth})()")
                    assert await page.evaluate('document.documentElement.scrollWidth<=innerWidth')
                    await page.screenshot(path=str(out / (language + '-mobile.png')), full_page=True)
                    await page.locator('#theme').click()
                    await page.screenshot(path=str(out / (language + '-dark-mobile.png')), full_page=True)
                    await page.locator('[data-report-section=money][role=tab]').click()
                    await page.emulate_media(media='print')
                    await page.evaluate("document.body.classList.add('print-finance')")
                    await expect(page.locator('#sidebar')).not_to_be_visible()
                    await expect(page.locator('#finance-report-form')).not_to_be_visible()
                    await expect(page.locator('.fin-print-table')).to_be_visible()
                    passed.append(language + ': relevant filters, mobile, dark and report-only printing')
                    await page.emulate_media(media='screen')
                    await page.evaluate("document.body.classList.remove('print-finance')")
                    await page.evaluate("liveCall('users.save',{name:'Report only QA',username:'reportonly',password:'0000',permissions:['reports.view'],active:true})")
                    await page.locator('#menu').click()
                    await page.locator('#demo-logout').click()
                    await expect(page.locator('#preview-login')).to_be_visible()
                    assert await page.evaluate('financeData===null')
                    await page.locator('#pages-worker').click()
                    await page.locator('#live-login [name="username"]').fill('reportonly')
                    await page.locator('#live-login [name="password"]').fill('0000')
                    await page.locator('#preview-enter').click()
                    await expect(page.locator('#preview-login')).to_be_hidden()
                    await expect(page.locator('#nav-workspace')).to_have_count(0)
                    await page.evaluate("navigate('reports')")
                    await expect(page.locator('[data-report-mode="other"]')).to_have_count(0)
                    await page.locator('#finance-report-form button.primary').click()
                    await page.wait_for_function('financeData!==null')
                    assert await page.evaluate('financeData.items.every(r=>!r.label&&!r.party&&!r.student_id&&!r.source_id)')
                    passed.append(language + ': scoped report user sees totals without private details or empty tools')
                    await page.close()
                assert not errors, errors
            finally:
                await browser.close()
                server.shutdown()
    (out / 'results.json').write_text(json.dumps({'passed': passed, 'errors': errors}, ensure_ascii=False, indent=2), encoding='utf-8')
    print('PERIOD_REPORTS_BROWSER_OK', len(passed), errors)


if __name__ == '__main__': asyncio.run(main())
