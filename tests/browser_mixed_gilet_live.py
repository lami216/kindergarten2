"""Actual protected windows, partial garment correction, PDF/XLSX and print."""
import asyncio,json,os,sys,uuid
import re
from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from playwright.async_api import async_playwright,expect


async def main(origin,username,password,dsn):
    from tests.qa_live_pages import guard
    guard(dsn)
    output=Path(__file__).resolve().parents[1]/'visual-evidence'/'mixed-gilet';output.mkdir(parents=True,exist_ok=True)
    errors=[]
    async with async_playwright() as pw:
        browser=await pw.chromium.launch(channel='chrome' if sys.platform=='win32' else None)
        try:
            for lang,width in [('ar',1360),('fr',390),('ar',390),('fr',1360)]:
                context=await browser.new_context(ignore_https_errors=True,viewport=dict(width=width,height=1000))
                page=await context.new_page();page.on('pageerror',lambda e:errors.append(str(e)))
                await page.goto(origin,wait_until='networkidle')
                if lang=='fr':
                    await page.locator('[data-language=fr]').click();await page.wait_for_load_state('networkidle')
                await page.locator('#live-login [name=username]').fill(username)
                await page.locator('#live-login [name=password]').fill(password)
                await page.locator('#preview-enter').click();await expect(page.locator('#preview-login')).to_be_hidden()
                async def api(op,data=None):return await page.evaluate('async([op,data])=>liveCall(op,data===null?undefined:data)',[op,data])
                state=await api('snapshot');day=date.fromisoformat(state['today']);start=day.replace(day=1)
                before=await api('reports.finance',dict(start=str(start),end=str(day)))
                prices=state['prices'];cash=sum(int(prices[k]) for k in ('registration','gilet','tuition'))
                retained=cash-int(prices['gilet']);department=next(x for x in state['departments'] if not x['cancelled'])
                name=f'Mixed receipt QA {lang} {width} '+uuid.uuid4().hex[:5]
                child=await api('students.create',dict(department_id=department['id'],name=name,birth='2022-01-01',gender='ذكر',
                    mother='Synthetic guardian',motherPhone='22000000',mode='pay',amount=str(cash),registration=True,gilet=True,request_token=str(uuid.uuid4())))
                sid=child['id'];await page.evaluate('async()=>{await liveLoad();render()}')
                state=await api('snapshot');rid=next(r['id'] for r in state['receipts'] if r['student']==sid)

                # A real information-only window must never inherit a warning.
                await page.evaluate('()=>summaryWindow(I18N.tr("ui.d06ec6b8635e"))')
                await expect(page.locator('#record-dialog')).to_be_visible();await expect(page.locator('.discard-box')).to_have_count(0)
                await page.locator('#workspace-close').click();await expect(page.locator('#record-dialog')).not_to_be_visible()

                # Exercise the actual student editor, not just injected miniature forms.
                await page.evaluate('sid=>{profileId=sid;page="students";render();}',sid)
                await page.locator('#open-profile-edit').click()
                field=page.locator('#profile-edit-child [name=address]');original=await field.input_value()
                await field.fill('Synthetic unsaved address');await expect(page.locator('.discard-box')).to_have_count(0)
                await page.locator('#profile-edit-close').click();await expect(page.locator('#profile-edit-dialog .discard-box')).to_be_visible()
                await page.locator('#profile-edit-dialog [data-keep]').click();await expect(field).to_have_value('Synthetic unsaved address')
                await field.fill(original);await page.locator('#profile-edit-close').click();await expect(page.locator('#profile-edit-dialog')).not_to_be_visible()

                await page.evaluate('id=>openReceipt(id)',rid)
                await page.locator('[data-cancel-gilet-item]').click();await expect(page.locator('#mixed-gilet-cancel-form')).to_be_visible()
                await expect(page.locator('.discard-box')).to_have_count(0)
                form=page.locator('#mixed-gilet-cancel-form')
                await form.locator('[name=reason]').fill('Correct only the garment item')
                await page.locator('#workspace-close').click();await expect(page.locator('.discard-box')).to_have_count(1)
                await page.locator('[data-keep]').click();await expect(form.locator('[name=reason]')).to_have_value('Correct only the garment item')
                await form.locator('[name=confirmed]').check()
                if lang=='fr' and width==1360:
                    await page.evaluate('()=>document.body.classList.add("dark")')
                await page.locator('#record-dialog').screenshot(path=str(output/f'{lang}-{width}-confirmation.png'))
                await form.locator('button.primary').click();await expect(page.locator('#record-dialog')).not_to_be_visible()
                await expect(page.locator('#receipt-page')).to_be_visible()
                active=await page.evaluate('receipts.find(r=>r.id===receiptPageId)')
                assert active['amount']==retained and active['gilet']==0 and active['registration']==prices['registration'] and active['monthly']==prices['tuition']
                assert active['date']==str(day) and active['corrects_id']==rid
                assert await page.locator('[data-cancel-gilet-item]').count()==0
                state=await api('snapshot');s=next(s for s in state['students'] if s['id']==sid)
                assert s['account']['due']==0 and s['account']['paid']==retained and s['account']['gilet_cancelled']
                for receipt in (rid,active['id']):
                    pdf=await page.request.get(origin+'/api/receipt/'+str(receipt))
                    assert pdf.status==200 and (await pdf.body()).startswith(b'%PDF')
                await page.locator('[data-receipt-chain="'+str(rid)+'"]').click()
                assert await page.evaluate('receipts.find(r=>r.id===receiptPageId).cancelled')
                await page.locator('[data-receipt-chain="'+str(active['id'])+'"]').click()

                # Actual reports and printed/exported details reconcile with the net receipt.
                if width<800:await page.locator('#menu').click()
                await page.locator('#nav-reports').click()
                await page.locator('#finance-report-form button.primary').click()
                await page.wait_for_function('()=>financeData!==null&&!financeLoading')
                report=await page.evaluate('financeData')
                assert report['summary']['income']==before['summary']['income']+retained
                assert report['summary']['gilet']==before['summary']['gilet']
                assert report['summary']['registration']==before['summary']['registration']+prices['registration']
                assert report['summary']['monthly']==before['summary']['monthly']+prices['tuition']
                await page.locator('[data-report-section=money][role=tab]').click();await page.locator('#fin-query').fill(name)
                await expect(page.locator('.fin-screen-table tbody tr')).to_have_count(1)
                assert await page.evaluate('financeRows()[0].amount')==retained
                async with page.expect_download() as download:
                    await page.locator('#fin-xlsx').click()
                path=output/f'{lang}-{width}.xlsx';await (await download.value).save_as(path)
                sheet=load_workbook(path).worksheets[1]
                assert sheet.max_row==2 and any(cell.value==name for row in sheet for cell in row)
                await page.evaluate('()=>{window.__mixedPrint=0;window.print=()=>window.__mixedPrint++;}')
                await page.locator('#fin-print').click();assert await page.evaluate('window.__mixedPrint')==1
                await page.emulate_media(media='print');await page.pdf(path=str(output/f'{lang}-{width}-report.pdf'),format='A4')
                printed=' '.join(p.extract_text() or '' for p in PdfReader(output/f'{lang}-{width}-report.pdf').pages)
                assert re.sub(r'\s','',name) in re.sub(r'\s','',printed)
                assert str(retained) in re.sub(r'[\s,\u202f]','',printed)
                await page.emulate_media(media='screen');await page.evaluate('()=>window.dispatchEvent(new Event("afterprint"))')
                assert await page.locator('html').get_attribute('dir')==('rtl' if lang=='ar' else 'ltr')
                assert not await page.evaluate('document.documentElement.scrollWidth>innerWidth+1')
                await page.evaluate('()=>navigate("history")')
                await page.locator('#history-type').select_option('monthly_gilet_cancellations')
                await expect(page.locator('#history-rows')).to_contain_text(name)
                await context.close()
        finally:await browser.close()
    assert not errors,errors
    from tests.browser_unsaved_windows import main as window_cases
    await window_cases(origin)
    print('MIXED_GILET_LIVE_OK: 4 actual bilingual desktop/mobile journeys; receipts, allocations, XLSX, print and real protected windows verified')


if __name__=='__main__':
    state=json.loads((Path(os.environ['KG_LIVE_QA_DIR'])/'state.json').read_text())
    asyncio.run(main(state['origin'],state['username'],state['password'],state['dsn']))
