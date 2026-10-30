"""Real HTTP/PG regressions; synthetic data and isolated schema only."""
from datetime import timedelta
import uuid
import pytest
import monthly_finance as F
import access_control as A
from tests.test_access_control import access, PASSWORD
from tests.test_live_pages_http import browser, login, post, intake


def child(client, session, **values):
    r=post(client,session,'students.create',intake(**values))
    assert r.status_code==200,r.text
    return r.json()['id']


def preview(client, session, sid, kind='gilet'):
    r=post(client,session,kind+'.preview',dict(student=sid))
    assert r.status_code==200,r.text
    return r.json()


def sale(client,session,sid,**changes):
    p=preview(client,session,sid)
    data=dict(student=sid,price=p['price'],tariff_revision=p['tariff_revision'],
        date=str(F.today()),confirmed=True,request_token=str(uuid.uuid4()))
    data.update(changes)
    return post(client,session,'gilet.sell',data),data


def exemption(client,session,sid,**changes):
    p=preview(client,session,sid,'exemption')
    data=dict(student=sid,revision=p['revision'],ledger_revision=p['ledger_revision'],
        due=p['due'],reason='Approved school exemption',confirmed=True,request_token=str(uuid.uuid4()))
    data.update(changes)
    return post(client,session,'exemption.apply',data),data


def test_independent_sales_replay_frozen_price_cancel_and_reports(browser):
    client,conn=browser;session=login(client);sid=child(client,session)
    before=F.projection(conn,sid)
    r,data=sale(client,session,sid);assert r.status_code==200,r.text
    first=r.json()['id']
    same=post(client,session,'gilet.sell',data)
    assert same.status_code==200 and same.json()['id']==first and not same.json()['created']
    assert post(client,session,'gilet.sell',{**data,'date':str(F.today()-timedelta(days=1))}).status_code==422
    second,_=sale(client,session,sid);assert second.status_code==200 and second.json()['id']!=first
    after=F.projection(conn,sid)
    assert before['account']['start_date']==after['account']['start_date']
    assert before['periods']==after['periods'] and before['due_now']==after['due_now']
    report=post(client,session,'reports.finance',dict(start=str(F.today()),end=str(F.today()))).json()
    assert report['summary']['gilet']==report['summary']['income']==4000
    assert F.financial_report(conn,F.today(),F.today())['gilet_cash']==4000
    F.update_tariffs(1000,2500,3000,'owner')
    from receipt_history import receipt_detail
    assert receipt_detail(conn,first)['items'][0]['stored_fee_price']==2000
    F.cancel_receipt(first,'Wrong sale',actor='owner')
    assert F.projection(conn,sid)['cash_paid']==2000
    report=post(client,session,'reports.finance',dict(start=str(F.today()),end=str(F.today()))).json()
    assert report['summary']['gilet']==report['summary']['income']==2000
    p=preview(client,session,sid)
    cancelled=next(x for x in p['purchases'] if x.get('receipt')==first)
    assert cancelled['delivered'] and cancelled['cancelled']
    from data_export import build_export,verify_export
    archive,_=build_export();manifest=verify_export(archive)
    assert next(t for t in manifest['tables'] if t['name']=='gilet_sales')['rows']==2
    assert next(t for t in manifest['tables'] if t['name']=='gilet_deliveries')['rows']==2


def test_exempt_sale_and_legacy_paid_delivery_without_cash(browser):
    client,conn=browser;session=login(client)
    sid=child(client,session,mode='exempt')
    r,_=sale(client,session,sid);assert r.status_code==200,r.text
    assert F.projection(conn,sid)['exempt']
    old=child(client,session,mode='pay',gilet=True,registration=False,amount='2000')
    p=preview(client,session,old);legacy=next(x for x in p['purchases'] if x['kind']=='legacy')
    assert not legacy['delivered'] and legacy['delivery_unknown']
    before=F.projection(conn,old)['cash_paid']
    data=dict(student=old,fee=legacy['id'],ledger_revision=p['ledger_revision'],
        reason='Handed to parent',confirmed=True,request_token=str(uuid.uuid4()))
    r=post(client,session,'gilet.deliver',data);assert r.status_code==200,r.text
    assert post(client,session,'gilet.deliver',data).status_code==200
    assert F.projection(conn,old)['cash_paid']==before
    assert preview(client,session,old)['purchases'][0]['delivered']


def test_unpaid_legacy_fee_uses_existing_settlement(browser):
    import student_debt as D
    client,conn=browser;session=login(client)
    sid=child(client,session,mode='pay',gilet=True,registration=False,amount='0')
    p=preview(client,session,sid);item=p['purchases'][0]
    assert item['due']==2000 and not item['delivered']
    invalid=dict(student=sid,fee=item['id'],ledger_revision=p['ledger_revision'],reason='No cash',confirmed=True,request_token=str(uuid.uuid4()))
    assert post(client,session,'gilet.deliver',invalid).status_code==422
    D.pay_unpaid_fee(sid,'gilet',actor='owner')
    assert preview(client,session,sid)['purchases'][0]['due']==0
    assert len(preview(client,session,sid)['purchases'])==1


def test_cancelled_payment_exemption_is_atomic_noncash_and_replayable(browser):
    client,conn=browser;session=login(client)
    sid=child(client,session,mode='pay',gilet=True,registration=True,amount='6000')
    paid=F.projection(conn,sid)['receipts'][0]['receipt_id']
    F.cancel_receipt(paid,'Mistaken receipt',actor='owner')
    before=F.projection(conn,sid);assert before['due_now']>0
    r,data=exemption(client,session,sid);assert r.status_code==200,r.text
    assert r.json()['waived']==before['due_now']
    p=F.projection(conn,sid)
    assert p['exempt'] and p['due_now']==0 and p['cash_paid']==before['cash_paid']
    assert post(client,session,'exemption.apply',data).json()['created'] is False
    import account_records as AR
    AR.change_account(sid,actor='owner',reason='Return to paid',start_date=p['account']['start_date'],status='regular')
    assert F.projection(conn,sid)['due_now']==0


def test_exemption_stale_preview_missing_permission_and_failure_rollback(browser,monkeypatch):
    client,conn=browser;session=login(client);sid=child(client,session)
    p=preview(client,session,sid,'exemption')
    data=dict(student=sid,revision=p['revision'],ledger_revision=p['ledger_revision'],due=p['due'],
        reason='Reviewed',confirmed=True,request_token=str(uuid.uuid4()))
    F.take_payment(sid,100,actor='owner')
    assert post(client,session,'exemption.apply',data).status_code==422
    assert not F.projection(conn,sid)['exempt']
    A.save_user(employee_name='No waiver rights',name='limited_exemption',password=PASSWORD,permissions=['accounts.view','accounts.edit'])
    limited=login(client,'limited_exemption')
    r,_=exemption(client,limited,sid);assert r.status_code==403,r.text
    session=login(client)
    original=F._audit
    def fail(c,student,receipt,action,*args,**kw):
        if action=='account_input_corrected':raise ValueError('Injected failure')
        return original(c,student,receipt,action,*args,**kw)
    monkeypatch.setattr(F,'_audit',fail)
    before=F.projection(conn,sid)
    r,_=exemption(client,session,sid);assert r.status_code==422,r.text
    after=F.projection(conn,sid)
    assert after['due_now']==before['due_now'] and not after['exempt']
    with conn.cursor() as c:
        c.execute('SELECT COUNT(*) n FROM monthly_waiver_batches WHERE student_id=%s',(sid,))
        assert c.fetchone()['n']==0


def test_no_debt_no_fake_waiver_and_archived_rejected(browser):
    import student_debt as D
    client,conn=browser;session=login(client)
    sid=child(client,session,mode='pay',registration=True,amount='4000')
    r,_=exemption(client,session,sid);assert r.status_code==200 and r.json()['waived']==0
    with conn.cursor() as c:
        c.execute('SELECT COUNT(*) n FROM monthly_waiver_batches WHERE student_id=%s',(sid,));assert c.fetchone()['n']==0
    D.archive_student(sid,actor='owner')
    r,_=exemption(client,session,sid);assert r.status_code==422
    r,_=sale(client,session,sid);assert r.status_code==422


def test_sale_date_correction_pdf_excel_and_price_revalidation(browser):
    from io import BytesIO
    from openpyxl import load_workbook
    from receipt_history import receipt_detail,receipt_impact
    client,conn=browser;session=login(client);sid=child(client,session,mode='exempt')
    p=preview(client,session,sid)
    F.update_tariffs(1000,2500,3000,'owner')
    r=post(client,session,'gilet.sell',dict(student=sid,price=p['price'],tariff_revision=p['tariff_revision'],
        date=str(F.today()),confirmed=True,request_token=str(uuid.uuid4())))
    assert r.status_code==422
    r,_=sale(client,session,sid);assert r.status_code==200,r.text
    original=r.json()['id'];day=F.today()-timedelta(days=1)
    frozen=receipt_detail(conn,original)
    impact=receipt_impact(conn,original,new_date=str(day))
    assert impact['cash_delta']==0 and impact['projected_due']==0
    rid,_=F.correct_receipt(original,2500,'Correct cash day',received_on=day,actor='owner',expected_revision=impact['revision'])
    original_items=receipt_detail(conn,original)['items']
    assert all(not item['active'] for item in original_items)
    assert [{k:v for k,v in item.items() if k!='active'} for item in original_items]==[{k:v for k,v in item.items() if k!='active'} for item in frozen['items']]
    assert receipt_detail(conn,rid)['items'][0]['stored_fee_price']==2500
    assert client.get('/api/receipt/'+str(rid)).content.startswith(b'%PDF')
    report=post(client,session,'reports.finance',dict(start=str(day),end=str(day))).json()
    assert report['summary']['gilet']==report['summary']['income']==2500
    exported=post(client,session,'reports.finance.export',dict(start=str(day),end=str(day),revision=report['revision']))
    assert exported.status_code==200,exported.text
    book=load_workbook(BytesIO(exported.content));assert len(book.sheetnames)==6
    assert len(preview(client,session,sid)['purchases'])==1
    with pytest.raises(ValueError):F.correct_receipt(rid,2400,'Wrong amount',actor='owner')
    assert F.projection(conn,sid)['cash_paid']==2500


def test_sale_atomic_rollback_no_orphan_garment_or_cash(browser,monkeypatch):
    client,conn=browser;session=login(client);sid=child(client,session)
    original=F._audit
    def fail(c,student,receipt,action,*args,**kw):
        if action=='gilet_sold':raise ValueError('Injected sale failure')
        return original(c,student,receipt,action,*args,**kw)
    monkeypatch.setattr(F,'_audit',fail)
    r,_=sale(client,session,sid);assert r.status_code==422
    with conn.cursor() as c:
        for table in ('gilet_sales','gilet_deliveries','monthly_receipts'):
            c.execute('SELECT COUNT(*) n FROM '+table);assert c.fetchone()['n']==0


def test_tool_route_cannot_bypass_debt_confirmation(browser):
    import records as R
    client,conn=browser;session=login(client);sid=child(client,session)
    with conn.cursor() as c:
        c.execute('SELECT * FROM monthly_accounts WHERE student_id=%s',(sid,));row=dict(c.fetchone())
    payload=dict(table='monthly_accounts',id=sid,mode='edit',revision=R.revision(row),
        values=dict(start_date=str(row['start_date']),status='exempt'),reason='Unreviewed debt',confirmed=True,request_token=str(uuid.uuid4()))
    assert post(client,session,'tools.change',payload).status_code==422
    p=preview(client,session,sid,'exemption')
    result=post(client,session,'tools.change',{**payload,'due':p['due'],'ledger_revision':p['ledger_revision'],'revision':p['revision']})
    assert result.status_code==200,result.text
    assert F.projection(conn,sid)['due_now']==0


@pytest.mark.parametrize('same_token',[True,False])
def test_concurrent_garment_requests_same_and_distinct_identities(browser,same_token):
    from concurrent.futures import ThreadPoolExecutor
    import gilet_sales as G
    import web_runtime as W
    client,conn=browser;session=login(client);sid=child(client,session)
    p=preview(client,session,sid);auth=A.login('owner',PASSWORD)
    key=str(uuid.uuid4())
    def run(token):
        with W.request_context(conn.dsn,{'auth_session':auth}):
            return G.sell(sid,actor='owner',price=p['price'],tariff_revision=p['tariff_revision'],received_on=str(F.today()),confirmed=True,token=token)
    with ThreadPoolExecutor(max_workers=2) as pool:
        result=list(pool.map(run,[key,key if same_token else str(uuid.uuid4())]))
    assert len({r[0] for r in result})==(1 if same_token else 2)
    assert sum(r[1] for r in result)==(1 if same_token else 2)
    assert F.projection(conn,sid)['cash_paid']==(2000 if same_token else 4000)


@pytest.mark.parametrize('endpoint',['gilet.sell','gilet.deliver','exemption.apply'])
def test_http_origin_csrf_and_permission_guards(browser,endpoint):
    from tests.test_live_pages_http import ORIGIN
    client,conn=browser;session=login(client);sid=child(client,session)
    assert client.post('/api/'+endpoint,json=dict(student=sid),headers={'Origin':ORIGIN}).status_code==403
    assert client.post('/api/'+endpoint,json=dict(student=sid),headers={'Origin':'https://other.test','X-CSRF-Token':session['csrf']}).status_code==403
    A.save_user(employee_name='Student reader',name='garment_reader',password=PASSWORD,permissions=['students.view'])
    limited=login(client,'garment_reader')
    if endpoint=='gilet.sell':
        response=post(client,limited,endpoint,dict(student=sid,price=2000,tariff_revision='x'*64,date=str(F.today()),confirmed=True,request_token=str(uuid.uuid4())))
    elif endpoint=='gilet.deliver':
        response=post(client,limited,endpoint,dict(student=sid,fee=1,ledger_revision=0,reason='Test',confirmed=True,request_token=str(uuid.uuid4())))
    else:
        response=post(client,limited,endpoint,dict(student=sid,revision='x'*64,ledger_revision=0,due=1000,reason='Test',confirmed=True,request_token=str(uuid.uuid4())))
    assert response.status_code==403,response.text


def test_upgrade_from_v16_preserves_legacy_rows_and_unknown_delivery(browser):
    import db,schema_migrations as M
    client,conn=browser;session=login(client)
    sid=child(client,session,mode='pay',registration=True,gilet=True,amount='6000')
    tables=('students','parents','monthly_accounts','monthly_fees','monthly_receipts','monthly_allocations',
        'monthly_allocation_snapshots','app_users','app_sessions','record_changes')
    with conn.cursor() as c:
        # Reconstruct the prior schema shape only inside this disposable schema.
        c.execute('ALTER TABLE monthly_allocations DROP CONSTRAINT monthly_allocation_target')
        c.execute('ALTER TABLE monthly_allocations DROP COLUMN gilet_sale_id CASCADE')
        c.execute('ALTER TABLE monthly_allocations ADD CONSTRAINT monthly_allocations_check CHECK((fee_id IS NULL)<>(period_id IS NULL))')
        c.execute('DROP TABLE gilet_deliveries; DROP TABLE gilet_sales; DROP TABLE account_status_requests')
        c.execute('DELETE FROM app_schema_migrations WHERE version=17')
        c.execute('INSERT INTO app_schema_migrations(version) VALUES(16) ON CONFLICT DO NOTHING')
        before={}
        for table in tables:
            c.execute('SELECT * FROM '+table);before[table]=[dict(r) for r in c.fetchall()]
    transaction=db.new_transaction_connection()
    try:assert M.upgrade(transaction)
    finally:transaction.close()
    with conn.cursor() as c:
        for table in tables:
            c.execute('SELECT * FROM '+table);after=[dict(r) for r in c.fetchall()]
            if table=='monthly_allocations':after=[{k:v for k,v in r.items() if k!='gilet_sale_id'} for r in after]
            assert after==before[table],table
        c.execute('SELECT COUNT(*) n FROM gilet_deliveries');assert c.fetchone()['n']==0
        c.execute('SELECT MAX(version) v FROM app_schema_migrations');assert c.fetchone()['v']==17
    assert preview(client,session,sid)['purchases'][0]['delivery_unknown']


def test_cancelled_legacy_receipt_cannot_authorize_delivery(browser):
    client,conn=browser;session=login(client)
    sid=child(client,session,mode='pay',gilet=True,registration=False,amount='2000')
    old=preview(client,session,sid);fee=old['purchases'][0]['id']
    rid=F.projection(conn,sid)['receipts'][0]['receipt_id']
    F.cancel_receipt(rid,'Never received payment',actor='owner')
    data=dict(student=sid,fee=fee,reason='Must not deliver',confirmed=True,request_token=str(uuid.uuid4()),ledger_revision=F.receipt_revision(conn,sid))
    assert post(client,session,'gilet.deliver',data).status_code==422


def test_existing_exemption_is_noop_and_successful_receipts_survive_waiver(browser):
    client,conn=browser;session=login(client)
    sid=child(client,session,mode='pay',registration=True,gilet=True,amount='3000',start_date=str(F.today()-timedelta(days=40)))
    original=F.projection(conn,sid);paid=original['cash_paid'];assert paid==3000 and original['due_now']>0
    r,_=exemption(client,session,sid);assert r.status_code==200
    p=F.projection(conn,sid);assert p['cash_paid']==paid and p['due_now']==0
    with conn.cursor() as c:
        c.execute('SELECT COUNT(*) n FROM monthly_audit WHERE student_id=%s',(sid,));n=c.fetchone()['n']
    r,_=exemption(client,session,sid);assert r.status_code==200 and not r.json()['created']
    with conn.cursor() as c:
        c.execute('SELECT COUNT(*) n FROM monthly_audit WHERE student_id=%s',(sid,));assert c.fetchone()['n']==n


def test_payment_racing_exemption_never_waives_unreviewed_debt(browser):
    import account_records as AR
    import web_runtime as W
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    client,conn=browser;session=login(client);sid=child(client,session)
    p=preview(client,session,sid,'exemption');auth=A.login('owner',PASSWORD);barrier=Barrier(2)
    def run(kind):
        with W.request_context(conn.dsn,{'auth_session':auth}):
            barrier.wait(timeout=15)
            try:
                if kind=='pay':F.take_payment(sid,100,actor='owner')
                else:AR.change_account(sid,actor='owner',reason='Concurrent exemption',status='exempt',
                    expected=p['revision'],expected_due=p['due'],expected_ledger=p['ledger_revision'],confirmed=True,token=str(uuid.uuid4()))
                return kind,True
            except ValueError:return kind,False
    with ThreadPoolExecutor(max_workers=2) as pool:result=dict(pool.map(run,['pay','exempt']))
    assert sum(result.values())==1
    after=F.projection(conn,sid)
    assert after['cash_paid']==(100 if result['pay'] else 0)
    assert after['exempt']==result['exempt']
    if result['exempt']:assert after['due_now']==0
