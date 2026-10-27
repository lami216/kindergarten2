"""Bilingual period school report, shared-family modal and protected drafts."""
import asyncio
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import json
from pathlib import Path
import sys
import tempfile
import threading

from openpyxl import load_workbook
from playwright.async_api import async_playwright, expect
from scripts.build_pages_preview import build


async def main():
    out=Path(__file__).resolve().parents[1]/'visual-evidence'/'school-period'
    out.mkdir(parents=True,exist_ok=True);results=[];errors=[]
    with tempfile.TemporaryDirectory() as folder:
        build(Path(folder))
        server=ThreadingHTTPServer(('127.0.0.1',0),partial(SimpleHTTPRequestHandler,directory=folder))
        threading.Thread(target=server.serve_forever,daemon=True).start()
        async with async_playwright() as pw:
            browser=await pw.chromium.launch(channel='chrome' if sys.platform=='win32' else None)
            try:
                for language in ('ar','fr'):
                    page=await browser.new_page(viewport={'width':1440,'height':1050})
                    page.on('pageerror',lambda error:errors.append(str(error)))
                    await page.goto(f'http://127.0.0.1:{server.server_port}/',wait_until='networkidle')
                    if language=='fr':await page.locator('[data-language="fr"]').click()
                    await page.locator('#preview-enter').click()
                    await expect(page.locator('#preview-login')).to_be_hidden()
                    await page.locator('#nav-reports').click()
                    await page.locator('[data-fin-period="last7"]').click()
                    await expect(page.locator('.school-cards')).to_be_visible()
                    await expect(page.locator('.fin-screen-table')).to_have_count(0)
                    await expect(page.locator('.school-screen-table')).to_have_count(0)
                    await page.screenshot(path=str(out/(language+'-overview.png')),full_page=True)
                    data=await page.evaluate('financeData')
                    a=data['school']['attendance'];s=a['summary']
                    assert data['comparison']['end']<data['start']
                    await page.locator('[data-report-section=overview][role=tab]').focus()
                    await page.keyboard.press('End')
                    await expect(page.locator('[data-report-section=attendance][role=tab]')).to_have_attribute('aria-selected','true')
                    await page.keyboard.press('Home')
                    await expect(page.locator('[data-report-section=overview][role=tab]')).to_have_attribute('aria-selected','true')
                    await page.locator('[data-report-section=money][role=tab]').click()
                    await page.locator('#fin-query').fill('unmatched hidden filter')
                    await page.locator('[data-report-section=overview][role=tab]').click()
                    assert s['total']==s['present']+s['absent']+s['excused'] and s['total']>0
                    assert s['attendance_rate']==round(s['present']*100/s['total'],1)
                    await page.locator('[data-report-section=attendance][role=tab]').click()
                    await page.locator('#school-status').select_option('absence')
                    expected=sum(row['absent']+row['excused']>0 for row in a['items'])
                    await expect(page.locator('.school-screen-table tbody tr')).to_have_count(expected)
                    await page.locator('#school-status').select_option('all')
                    await page.locator('#school-query').fill('(')
                    await expect(page.locator('.school-screen-table tbody tr')).to_have_count(0)
                    await page.locator('#school-query').fill('')
                    await page.locator('[data-school-tab="new"]').click()
                    await expect(page.locator('.school-screen-table tbody tr')).to_have_count(data['school']['new_count'])
                    async with page.expect_download() as event:await page.locator('#fin-xlsx').click()
                    download=await event.value;filename=out/(language+'.xlsx');await download.save_as(filename)
                    book=load_workbook(filename)
                    assert len(book.sheetnames)==6
                    assert book.worksheets[1].max_row==len(data['items'])+1
                    await page.locator('.school-period').screenshot(path=str(out/(language+'-desktop.png')))
                    await page.locator('[data-report-section=overview][role=tab]').click()
                    await page.set_viewport_size({'width':390,'height':900})
                    await page.wait_for_function("(()=>{const r=document.querySelector('#sidebar').getBoundingClientRect();return r.right<=0||r.left>=innerWidth})()")
                    await page.screenshot(path=str(out/(language+'-overview-mobile.png')),full_page=True)
                    await page.locator('[data-school-tab="attendance"]').click()
                    assert await page.evaluate('document.documentElement.scrollWidth<=innerWidth+1')
                    await page.screenshot(path=str(out/(language+'-mobile.png')),full_page=True)
                    await page.locator('#theme').click()
                    await page.locator('.school-period').screenshot(path=str(out/(language+'-dark.png')))
                    await page.evaluate("document.body.classList.add('print-finance')")
                    await page.emulate_media(media='print')
                    await expect(page.locator('.school-print-table')).to_be_visible()
                    await expect(page.locator('.report-print-heading')).to_contain_text(await page.evaluate('I18N.date(financeData.start)'))
                    await expect(page.locator('.school-screen-table')).to_be_hidden()
                    await page.emulate_media(media='screen')
                    await page.evaluate("document.body.classList.remove('print-finance')")
                    results.append(language+': period counts, weighted rates, literal search, workbook, mobile, dark and print')
                    # Two real preview children share one family, not just equal names.
                    ids=await page.evaluate('''async()=>{
                      const first=students.find(s=>!s.cancelled);
                      const sibling=await liveCall('students.create',{department_id:first.department_id||1,name:'Sibling QA',birth:'2022-01-01',gender:'ذكر',existing_parent_id:first.parent_id,mode:'pay',amount:'0',registration:false,gilet:false,request_token:crypto.randomUUID()});
                      await liveLoad();profileId=first.id;page='students';render();return {first:first.id,sibling:sibling.id};
                    }''')
                    await page.locator('#open-profile-edit').click()
                    await expect(page.locator('#profile-edit-dialog')).to_be_visible()
                    await page.locator('#profile-edit-child [name="motherPhone"]').fill('99887766')
                    await page.locator('#profile-edit-close').click()
                    await expect(page.locator('#profile-edit-dialog .discard-box')).to_be_visible()
                    await page.locator('#profile-edit-dialog [data-keep]').click()
                    await page.locator('#profile-edit-child .form-end button.primary').click()
                    await expect(page.locator('#profile-edit-dialog')).to_be_hidden()
                    assert await page.evaluate('(id)=>students.find(s=>s.id===id).motherPhone',ids['sibling'])=='99887766'
                    await page.locator('#open-profile-edit').click()
                    await page.locator('#profile-edit-child [name="motherPhone"]').fill('11111111')
                    await page.locator('#profile-edit-close').click()
                    await page.locator('#profile-edit-dialog [data-discard]').click()
                    assert await page.evaluate('(id)=>students.find(s=>s.id===id).motherPhone',ids['first'])=='99887766'
                    await page.locator('#open-profile-edit').click()
                    await page.locator('.profile-parents').screenshot(path=str(out/(language+'-parent-edit.png')))
                    await page.locator('#profile-edit-close').click()
                    # Parent editing stays hidden for a student editor without that grant.
                    await page.evaluate("async()=>{const r=await liveCall('users.save',{name:'Student editor',username:'schooleditor',password:'0000',permissions:['students.view','students.edit'],active:true});await liveCall('logout',{});await liveCall('login',{username:'schooleditor',password:'0000'});await liveEnter();}")
                    await page.evaluate("(id)=>{profileId=id;page='students';render()}",ids['first'])
                    await page.locator('#open-profile-edit').click()
                    await expect(page.locator('#profile-edit-child [name="motherPhone"]')).to_have_count(0)
                    results.append(language+': parent modal, shared sibling, unsaved discard and permission gating')
                    await page.close()
                assert not errors,errors
            finally:
                await browser.close();server.shutdown();server.server_close()
    (out/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'passed':len(results),'results':results},ensure_ascii=False))


if __name__=='__main__':asyncio.run(main())
