"""Literal XLSX values and identity-based payroll filters, including archives."""
from datetime import date, datetime
from decimal import Decimal
import io
import os
import uuid
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import load_workbook
import psycopg2
from psycopg2 import sql
import pytest
from streamlit.testing.v1 import AppTest

from sections import reports


@pytest.mark.parametrize('text', [
    '=1+1', '=SUM(1,2)', '=HYPERLINK("https://example.invalid","name")',
    '+123', '-123', '@SUM(A1:A2)', '\t=1+1', '\n=1+1', ' =1+1',
    '#N/A', '#DIV/0!', "'=1+1", 'طفل عادي', '00001234',
])
def test_report_excel_roundtrips_untrusted_text_without_formulas(text):
    frame=pd.DataFrame({'=1+1':[text], 'المبلغ':[123.5]})
    original=frame.copy(deep=True)
    payload=reports._to_excel(frame)
    for data_only in (False,True):
        book=load_workbook(io.BytesIO(payload),data_only=data_only)
        sheet=book['التقرير']
        assert sheet['A1'].value=='=1+1' and sheet['A1'].data_type=='s'
        assert sheet['A2'].value==text and sheet['A2'].data_type=='s'
        assert sheet['B2'].value==123.5 and sheet['B2'].data_type=='n'
        book.close()
    with ZipFile(io.BytesIO(payload)) as archive:
        for path in archive.namelist():
            if path.startswith('xl/worksheets/') and path.endswith('.xml'):
                root=ET.fromstring(archive.read(path))
                assert not root.findall('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f')
        assert not any(path.startswith('xl/externalLinks/') for path in archive.namelist())
    pd.testing.assert_frame_equal(frame,original)


def test_export_preserves_numeric_date_boolean_and_empty_types():
    payload=reports._to_excel(pd.DataFrame([{
        'اسم':'=1+1', 'مبلغ':Decimal('1200.50'), 'سالب':-10, 'صفر':0,
        'تاريخ':date(2026,10,18), 'مؤكد':True, 'فارغ':None,
    }]))
    sheet=load_workbook(io.BytesIO(payload),data_only=False)['التقرير']
    assert sheet['A2'].value=='=1+1'
    assert sheet['B2'].value==1200.5 and sheet['B2'].data_type=='n'
    assert sheet['C2'].value==-10 and sheet['D2'].value==0
    assert sheet['E2'].value==datetime(2026,10,18) and sheet['E2'].data_type=='d'
    assert sheet['F2'].value is True and sheet['F2'].data_type=='b'
    assert sheet['G2'].value is None


def test_teacher_choices_preserve_same_names_and_reused_archived_identities():
    frame=pd.DataFrame([
        dict(national_id='one',display_national_id='001',full_name='أمينة',archived_at=None),
        dict(national_id='two',display_national_id='002',full_name='أمينة',archived_at=None),
        dict(national_id='old1',display_national_id='001',full_name='أمينة',archived_at=date(2026,1,1)),
        dict(national_id='old2',display_national_id='001',full_name='أمينة',archived_at=date(2026,2,1)),
    ])
    labels=reports._teacher_labels(frame)
    assert set(labels)=={'one','two','old1','old2'} and len(set(labels.values()))==4
    assert labels['one']=='أمينة — 001' and labels['two']=='أمينة — 002'
    assert all('مؤرشفة' in labels[key] for key in ('old1','old2'))
    assert reports._teacher_labels(pd.DataFrame())=={}


@pytest.mark.skipif(not os.getenv('TEST_POSTGRES_DSN'),reason='Disposable PostgreSQL required')
def test_payroll_report_selects_each_same_named_teacher_independently(monkeypatch):
    schema='report_identity_'+uuid.uuid4().hex
    conn=psycopg2.connect(os.environ['TEST_POSTGRES_DSN']);conn.autocommit=True
    monkeypatch.setenv('REPORT_IDENTITY_SCHEMA',schema)
    try:
        with conn.cursor() as c:
            c.execute(sql.SQL('CREATE SCHEMA {}').format(sql.Identifier(schema)))
            c.execute(sql.SQL('SET search_path TO {}').format(sql.Identifier(schema)))
            c.execute('''CREATE TABLE teachers(national_id TEXT PRIMARY KEY,display_national_id TEXT,
                full_name TEXT,archived_at TIMESTAMPTZ);
                CREATE TABLE teacher_payments(payment_id INTEGER PRIMARY KEY,national_id TEXT REFERENCES teachers,
                salary_month TEXT,base_salary NUMERIC,bonus NUMERIC,deductions NUMERIC,amount NUMERIC,
                cash_amount NUMERIC,record_state TEXT,payment_date DATE)''')
            for i,(ident,display,amount,archived) in enumerate([
                ('stable-A','001',1000,None), ('stable-B','002',2400,None),
                ('archived-A','001',900,datetime(2025,1,1)),
            ],1):
                c.execute('INSERT INTO teachers VALUES(%s,%s,%s,%s)',(ident,display,'أمينة أحمد',archived))
                c.execute('''INSERT INTO teacher_payments VALUES(%s,%s,'2026-10',%s,0,0,%s,%s,'active','2026-10-18')''',
                          (i,ident,amount,amount,amount))
        app=AppTest.from_string('''
import os,streamlit as st,psycopg2
from psycopg2 import sql
from psycopg2.extras import RealDictCursor
from sections import reports
st.session_state.setdefault('report_type','💼 سجل الرواتب والمستحقات')
c=psycopg2.connect(os.environ['TEST_POSTGRES_DSN'],cursor_factory=RealDictCursor)
c.autocommit=True
with c.cursor() as q:q.execute(sql.SQL('SET search_path TO {}').format(sql.Identifier(os.environ['REPORT_IDENTITY_SCHEMA'])))
reports.render(c)
c.close()
''').run(timeout=30)
        assert not app.exception
        picker=app.selectbox(key='salary_report_teacher')
        assert len(picker.options)==4
        assert 'أمينة أحمد — 001' in picker.options and 'أمينة أحمد — 002' in picker.options
        assert any('مؤرشفة' in option for option in picker.options)
        for ident,display,amount in [('stable-A','001',1000),('stable-B','002',2400),('archived-A','001',900)]:
            app.selectbox(key='salary_report_teacher').set_value(ident).run(timeout=30)
            assert not app.exception
            frame=app.dataframe[0].value
            assert len(frame)==1 and frame.iloc[0]['هوية المعلم/ة']==display
            assert frame.iloc[0]['النقد المصروف']==amount
        app.selectbox(key='salary_report_teacher').set_value(None).run(timeout=30)
        assert not app.exception
        assert len(app.dataframe[0].value)==3
        assert sum(app.dataframe[0].value['النقد المصروف'])==4300
    finally:
        with conn.cursor() as c:
            c.execute('RESET search_path')
            c.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(schema)))
        conn.close()
