"""Period cash reports: totals, drill-down and export share one consistent ledger."""
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
import hashlib
from io import BytesIO
import json
import unicodedata

import access_control as A
import db
import localization as L
import monthly_finance as F

ZERO = Decimal(0)

# Corrected cash amounts, not invoice totals or payroll documents. Explicitly
# linked payroll expense rows are excluded exactly as in the main cash summary.
LEDGER = '''WITH allocations AS (
 SELECT a.receipt_id,
 COALESCE(SUM(a.amount) FILTER(WHERE f.kind='registration'),0) AS registration,
 COALESCE(SUM(a.amount) FILTER(WHERE f.kind='gilet'),0) AS gilet,
 COALESCE(SUM(a.amount) FILTER(WHERE a.period_id IS NOT NULL),0) AS monthly,
 COALESCE(SUM(a.amount) FILTER(WHERE p.starts>r.received_on),0) AS prepaid
 FROM monthly_allocations a JOIN monthly_receipts r USING(receipt_id)
 LEFT JOIN monthly_fees f USING(fee_id) LEFT JOIN monthly_periods p USING(period_id)
 WHERE a.active AND r.state='active' GROUP BY a.receipt_id
), amendments AS (
 SELECT event_id,SUM(delta) AS delta FROM expense_cash_amendments GROUP BY event_id
), ledger AS (
 SELECT 'receipt:'||r.receipt_id AS key,r.received_on AS on_date,'income' AS category,
 r.amount::numeric AS amount,s.full_name AS label,r.payer AS party,r.receipt_id AS source_id,
 r.student_id,COALESCE(a.registration,0) AS registration,COALESCE(a.gilet,0) AS gilet,
 COALESCE(a.monthly,0) AS monthly,COALESCE(a.prepaid,0) AS prepaid,false AS void_document
 FROM monthly_receipts r LEFT JOIN students s USING(student_id)
 LEFT JOIN allocations a USING(receipt_id) WHERE r.state='active'
 UNION ALL
 SELECT 'expense:'||x.event_id,x.event_date,'operational',x.amount+COALESCE(a.delta,0),
 e.category,e.payee,e.expense_id,NULL,0,0,0,0,false
 FROM expense_cash_events x JOIN expenses e USING(expense_id)
 LEFT JOIN amendments a USING(event_id)
 WHERE e.teacher_payment_id IS NULL AND e.teacher_advance_id IS NULL
 UNION ALL
 SELECT 'salary:'||p.payment_id,p.payment_date,'salary',COALESCE(p.cash_amount,p.amount),
 t.full_name,p.salary_month,p.payment_id,NULL,0,0,0,0,p.record_state='cancelled'
 FROM teacher_payments p LEFT JOIN teachers t USING(national_id)
 UNION ALL
 SELECT 'advance:'||p.advance_id,p.advance_date,'advance',COALESCE(p.cash_amount,p.amount),
 t.full_name,p.salary_month,p.advance_id,NULL,0,0,0,0,p.record_state='cancelled'
 FROM teacher_advances p LEFT JOIN teachers t USING(national_id)
), opening AS (
 SELECT COALESCE(SUM(CASE WHEN category='income' THEN amount ELSE -amount END),0) AS balance
 FROM ledger WHERE on_date<%(start)s AND amount>0
)
SELECT l.*,o.balance AS opening FROM opening o LEFT JOIN ledger l
 ON l.on_date BETWEEN %(start)s AND %(end)s AND l.amount>0
ORDER BY l.on_date DESC,l.key DESC'''


def label(key):
    return L.CATALOG['reports.' + key][L.language()]


def dates(data):
    from live_pages import text
    try:
        start = date.fromisoformat(text(data, 'start', 10, True))
        end = date.fromisoformat(text(data, 'end', 10, True))
    except ValueError:
        raise ValueError('راجع تاريخ البداية والنهاية.') from None
    if start > end or end > F.today():
        raise ValueError('راجع تاريخ البداية والنهاية.')
    return start, end


def school_period(cursor, start, end, permissions):
    """Enrollment dates and student-days from the same report snapshot as cash."""
    from subscription_dates import registration_sql
    cursor.execute('''SELECT s.student_id AS id,s.full_name AS name,s.archived_at IS NOT NULL AS archived,
        '''+registration_sql()+''' AS registered_on,d.name AS department
        FROM students s LEFT JOIN student_departments sd USING(student_id)
        LEFT JOIN monthly_accounts a USING(student_id)
        LEFT JOIN school_departments d USING(department_id) ORDER BY s.full_name,s.student_id''')
    students = [dict(row) for row in cursor.fetchall()]
    added = [dict(row, registered_on=str(row['registered_on'])) for row in students
             if row['registered_on'] is not None and start <= row['registered_on'] <= end]
    result = dict(new_count=len(added), new_students=added if permissions['students'] else [],
                  unknown_dates=sum(row['registered_on'] is None for row in students), attendance=None)
    if permissions['attendance']:
        from attendance import percentages
        cursor.execute('''SELECT a.person_id AS id,s.full_name AS name,s.archived_at IS NOT NULL AS archived,
            COUNT(*) FILTER(WHERE a.status='present') AS present,
            COUNT(*) FILTER(WHERE a.status='absent') AS absent,
            COUNT(*) FILTER(WHERE a.status='excused') AS excused
            FROM student_attendance a JOIN students s ON s.student_id=a.person_id
            WHERE a.cancelled_at IS NULL AND a.on_date BETWEEN %s AND %s
            GROUP BY a.person_id,s.full_name,s.archived_at ORDER BY s.full_name,a.person_id''', (start, end))
        rows = [dict(id=row['id'], name=row['name'], archived=row['archived'], **percentages(row))
                for row in cursor.fetchall()]
        cursor.execute('''SELECT COUNT(DISTINCT on_date) AS days FROM student_attendance
            WHERE cancelled_at IS NULL AND on_date BETWEEN %s AND %s''', (start, end))
        result['attendance'] = dict(summary=percentages({key: sum(row[key] for row in rows)
            for key in ('present', 'absent', 'excused')}), students_count=len(rows),
            recorded_days=cursor.fetchone()['days'], items=rows)
    return result


def previous_period(cursor, start, end, permissions):
    """Equal-length preceding range; aggregates only, inside the same snapshot."""
    try:
        previous_end = start - timedelta(days=1)
        previous_start = start - timedelta(days=(end-start).days+1)
    except OverflowError:
        return None
    cursor.execute('SELECT category,COALESCE(SUM(amount),0) AS amount FROM ('+LEDGER+
                   ') period_rows WHERE key IS NOT NULL GROUP BY category',
                   dict(start=previous_start,end=previous_end))
    amounts = {row['category']: row['amount'] for row in cursor.fetchall()}
    incoming = amounts.get('income', ZERO)
    outgoing = sum((amounts.get(k, ZERO) for k in ('operational','salary','advance')), ZERO)
    from subscription_dates import registration_sql
    cursor.execute('''SELECT COUNT(*) AS n FROM students s LEFT JOIN monthly_accounts a USING(student_id) WHERE
        '''+registration_sql()+''' BETWEEN %s AND %s''', (previous_start,previous_end))
    result = dict(start=str(previous_start),end=str(previous_end),new_count=cursor.fetchone()['n'],
                  summary=dict(income=incoming,outgoing=outgoing,net=incoming-outgoing),attendance=None)
    if permissions['attendance']:
        from attendance import percentages
        cursor.execute('''SELECT COUNT(*) FILTER(WHERE status='present') AS present,
            COUNT(*) FILTER(WHERE status='absent') AS absent,
            COUNT(*) FILTER(WHERE status='excused') AS excused
            FROM student_attendance WHERE cancelled_at IS NULL AND on_date BETWEEN %s AND %s''',
            (previous_start,previous_end))
        result['attendance'] = percentages(cursor.fetchone())
    return result


def read(data):
    start, end = dates(data)
    if not A.can('reports.view'):
        raise A.AccessDenied()
    # Schedule materialization is outside the read-only snapshot, never inside it.
    F._ensure_all_started(end)
    conn = db.new_transaction_connection()
    try:
        conn.rollback()
        conn.set_session(isolation_level='REPEATABLE READ', readonly=True)
        with conn.cursor() as c:
            profile = A.require(c, 'reports.view', lock=False)
            can = lambda area: A.allowed(profile, area + '.view')
            permissions = {area: can(area) for area in ('payments', 'expenses', 'salaries', 'students')}
            permissions['attendance'] = can('student_attendance')
            c.execute(LEDGER, {'start': start, 'end': end})
            raw = [dict(r) for r in c.fetchall()]
            opening = raw[0]['opening'] if raw else ZERO
            items = []
            for n, row in enumerate(raw):
                if row['key'] is None:
                    continue
                row.pop('opening')
                row['date'] = row.pop('on_date').isoformat()
                area = {'income': 'payments', 'operational': 'expenses', 'salary': 'salaries', 'advance': 'salaries'}[row['category']]
                if not permissions[area]:
                    row.update(key='restricted:' + str(n), label='', party='', source_id=None, student_id=None)
                items.append(row)
            # Debt is a stock at end date, not money received or paid in the period.
            accounts = F.projections(conn, end, materialize=False)
            student_total = sum((Decimal(p['due_now']) for p in accounts.values()), ZERO)
            student_debts = []
            if permissions['payments']:
                c.execute('SELECT student_id,full_name FROM students')
                names = {r['student_id']: r['full_name'] for r in c.fetchall()}
                student_debts = [dict(id=sid, name=names.get(sid, sid), amount=p['due_now'], archived=bool(p['account']['archived_at']))
                                 for sid, p in accounts.items() if p['due_now']]
                student_debts.sort(key=lambda r: str(r['id']))
            c.execute('''SELECT e.expense_id AS id,e.category AS name,e.payee,
                GREATEST(0,e.amount-COALESCE((SELECT SUM(x.amount+COALESCE(
                (SELECT SUM(am.delta) FROM expense_cash_amendments am WHERE am.event_id=x.event_id),0))
                FROM expense_cash_events x WHERE x.expense_id=e.expense_id AND x.event_date<=%s),0)) AS amount
                FROM expenses e WHERE e.expense_date<=%s AND e.cancelled_at IS NULL
                AND e.teacher_payment_id IS NULL AND e.teacher_advance_id IS NULL''', (end, end))
            supplier_debts = [dict(r) for r in c.fetchall() if r['amount']]
            supplier_debts.sort(key=lambda r: str(r['id']))
            supplier_total = sum((r['amount'] for r in supplier_debts), ZERO)
            summary = {key: sum((Decimal(r['amount']) for r in items if r['category'] == key), ZERO)
                       for key in ('income', 'operational', 'salary', 'advance')}
            summary['outgoing'] = sum((summary[k] for k in ('operational', 'salary', 'advance')), ZERO)
            summary['net'] = summary['income'] - summary['outgoing']
            summary.update(opening=opening, closing=opening + summary['net'])
            for key in ('registration', 'gilet', 'monthly', 'prepaid'):
                summary[key] = sum((Decimal(r[key]) for r in items if r['category'] == 'income'), ZERO)
            summary['other_income'] = summary['income'] - sum((summary[k] for k in ('registration', 'gilet', 'monthly')), ZERO)
            result = dict(start=str(start), end=str(end), summary=summary, items=items, permissions=permissions,
                          debt=dict(students_total=student_total, suppliers_total=supplier_total,
                                    students=student_debts, suppliers=supplier_debts if permissions['expenses'] else []))
            result['school'] = school_period(c, start, end, permissions)
            result['comparison'] = previous_period(c, start, end, permissions)
            result['revision'] = hashlib.sha256(json.dumps(result, default=str, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
            result['generated_at'] = datetime.now(timezone.utc).isoformat()
            return result
    finally:
        conn.close()


def filtered_items(report, data):
    from live_pages import text
    category = text(data, 'category', 20, default='all')
    if category not in ('all', 'income', 'outgoing', 'operational', 'salary', 'advance'):
        raise ValueError('التقرير غير متاح.')
    query = unicodedata.normalize('NFKC', text(data, 'query', 120)).lower()
    return [r for r in report['items']
            if (category == 'all' or r['category'] == category or category == 'outgoing' and r['category'] != 'income')
            and query in unicodedata.normalize('NFKC', ' '.join(str(r[k] or '') for k in ('label', 'party', 'source_id', 'date'))).lower()]


def export(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    report = read(data)
    if data.get('revision') != report['revision']:
        raise ValueError('تغيرت الحركات المالية؛ حدّث التقرير قبل تصديره.')
    rows = filtered_items(report, data)
    book = Workbook()
    summary = book.active
    summary.title = label('summarySheet')
    summary.append([label('title'), report['start'], report['end']])
    summary.append([label('generated'), report['generated_at']])
    summary.append([label('query'), data.get('query', ''), label(data.get('category', 'all'))])
    for key in ('income', 'outgoing', 'net', 'opening', 'closing', 'registration', 'gilet', 'monthly', 'prepaid', 'other_income', 'operational', 'salary', 'advance'):
        summary.append([label(key), report['summary'][key]])
    summary.append([label('studentsDebt'), report['debt']['students_total']])
    summary.append([label('suppliersDebt'), report['debt']['suppliers_total']])
    summary.append([label('filteredTotal'), sum((r['amount'] if r['category'] == 'income' else -r['amount'] for r in rows), ZERO)])
    summary.append([label('cashNote')])
    sheet = book.create_sheet(label('detailsSheet'))
    sheet.append([label(k) for k in ('date', 'type', 'description', 'party', 'incomingColumn', 'outgoingColumn', 'reference')])
    for row in rows:
        sheet.append([row['date'], label(row['category']+'Entry'), row['label'], row['party'],
                      row['amount'] if row['category'] == 'income' else ZERO,
                      row['amount'] if row['category'] != 'income' else ZERO, row['source_id']])
    debts = book.create_sheet(label('debtSheet'))
    debts.append([label('asOf'), report['end']])
    debts.append([label('type'), label('description'), label('amount')])
    for key, title in (('students', 'studentsDebt'), ('suppliers', 'suppliersDebt')):
        for row in report['debt'][key]:
            debts.append([label(title), row['name'], row['amount']])
    school = report['school']
    summary.append([label('newStudents'), school['new_count']])
    if report['permissions']['students']:
        added = book.create_sheet(label('newSheet'))
        added.append([label('student'), label('registeredOn'), label('department'), label('fileState')])
        for row in school['new_students']:
            added.append([row['name'], row['registered_on'], row['department'], label('archived') if row['archived'] else label('active')])
    if school['attendance'] is not None:
        attendance = school['attendance']
        for key in ('present', 'absent', 'excused', 'total', 'attendance_rate', 'absence_rate'):
            summary.append([label(key), attendance['summary'][key]])
        summary.append([label('recordedDays'), attendance['recorded_days']])
        summary.append([label('attendanceNote')])
        sheet = book.create_sheet(label('attendanceSheet'))
        sheet.append([label(key) for key in ('student', 'present', 'absent', 'excused', 'total', 'attendance_rate', 'absence_rate')])
        for row in attendance['items']:
            sheet.append([row['name'], *[row[key] for key in ('present', 'absent', 'excused', 'total', 'attendance_rate', 'absence_rate')]])
    previous = report['comparison']
    if previous is not None:
        sheet = book.create_sheet(label('comparisonSheet'))
        sheet.append([label('period'),report['start'],report['end']])
        sheet.append([label('previousRange'),previous['start'],previous['end']])
        sheet.append([label('metric'),label('currentRange'),label('previousRange'),label('difference')])
        for key in ('income','outgoing','net'):
            current, old = report['summary'][key], previous['summary'][key]
            sheet.append([label(key),current,old,current-old])
        current, old = school['new_count'], previous['new_count']
        sheet.append([label('newStudents'),current,old,current-old])
        if school['attendance'] is not None:
            current, old = school['attendance']['summary']['attendance_rate'], previous['attendance']['attendance_rate']
            sheet.append([label('attendance_rate'),current,old,current-old if current is not None and old is not None else None])
        sheet.append([label('comparisonNote')])
    for sheet in book:
        sheet.sheet_view.rightToLeft = L.language() == 'ar'
        sheet.freeze_panes = 'A2'
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
                if isinstance(cell.value, (int, float, Decimal)): cell.number_format = '#,##0.00'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='087F5B')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 26
    stream = BytesIO();book.save(stream)
    return stream.getvalue()
